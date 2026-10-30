#!/usr/bin/env python3
"""S16 F3–F4 – műhelyi dokumentáció egyetlen adatforrásból.

Bemenet:  ../10_F1_adatmodell/parts_v1.json  +  ../30_F2b_vasalat/vasalat_terv.json
Kimenet (ebben a mappában):
  S16_R15_szabasjegyzek.xlsx        szabásjegyzék, anyag-összesítő, vasalatlista, furatlista, nyitott kérdések
  S16_R15_alkatreszlapok.pdf        alkatrészenként egy A4 lap: kontúr, A/B oldali és élfuratok, élzárás, ütköző, szálirány, pecsét
  dxf_alkatreszek/<címke>.dxf      ugyanez DXF-ben (rétegek: KONTUR, FURAT_A, FURAT_B, FURAT_EL, ELZARAS, SZOVEG, MERET)
  S16_R15_cimkek.pdf                100×50 mm-es matricák, Code128 vonalkód = címke-ID (= MPR programnév alapja)
  mpr_tervezet/<címke>_A.mpr / _B.mpr   woodWOP MPR TERVEZET (furatok; kontúrmarás szerszámtábla nélkül nincs)
  dokumentacio_qa.md                visszaolvasó egyezés: adatmodell ↔ vasalatterv ↔ MPR ↔ Excel ↔ lapok

Gépi konvenció (TERVEZET, az asztalos minta-MPR-jéig): X = helyi U (hossz), Y = helyi V (szélesség), A oldal felül;
B oldali program: a darab az U tengely körül átfordítva (Y = W − v). Élfurat: E4 → XA=0 BM=XP, E2 → XA=L BM=XM,
E1 → YA=0 BM=YP, E3 → YA=W BM=YM; ZA = magasság a lap alsó síkjától (az A-oldali programban a B sík).
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
F1 = HERE.parent / "10_F1_adatmodell"
F2B = HERE.parent / "30_F2b_vasalat"
sys.path.insert(0, str(F1))
from s16_geom import AXIS_VEC, LocalFrame, load_json  # noqa: E402

TODAY = date.today().isoformat()
STATUS = "TERVI – NEM GYÁRTÁSI KIADÁS"
MPR_KEYS_VERT = ["XA", "YA", "BM", "TI", "DU", "AN", "MI", "AB", "F_", "EN", "KM"]
MPR_KEYS_HOR = ["XA", "YA", "ZA", "BM", "TI", "DU", "AN", "MI", "AB", "F_", "EN", "KM"]
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


# ---------------------------------------------------------------------------- adatelőkészítés

def load_all():
    doc = load_json(F1 / "parts_v1.json")
    rules = load_json(F1 / "gyartasi_szabalyok_v1.json")
    vas = load_json(F2B / "vasalat_terv.json")
    src = {p["id"]: p for p in load_json(F1 / doc["forras"]["mappa"] / "parts.json")}
    return doc, rules, vas, src


def classify_holes(doc, rules, vas):
    """Furatok alkatrészenként, helyi koordinátában, oldal/él szerint besorolva."""
    nez = rules["koordinata_konvencio"]["nezetek"]
    R = {r["id"]: r for r in doc["alkatreszek"]}
    out = defaultdict(list)
    for h in vas["furatok"]:
        r = R[h["alkatresz"]]
        fr = LocalFrame(r["befoglalo_globalis"], r["helyi_rendszer"]["A_normal"], nez)
        d = np.array(h["irany"], float)
        ld = np.array([d @ fr.U, d @ fr.V, d @ fr.N])
        u, v, n = fr.to_local(np.array(h["start"], float))
        L, W, T = fr.L, fr.W, fr.T
        k = int(np.argmax(np.abs(ld)))
        if k == 2:
            side = "A" if ld[2] < 0 else "B"
            el = None
        else:
            side = "EL"
            if k == 0:
                el = "E4" if ld[0] > 0 else "E2"
            else:
                el = "E1" if ld[1] > 0 else "E3"
        out[r["id"]].append({"u": round(float(u), 2), "v": round(float(v), 2), "n": round(float(n), 2), "oldal": side, "el": el,
                             "d": h["d"], "melyseg": h["melyseg"], "funkcio": h["funkcio"], "kapcsolat": h["kapcsolat"],
                             "L": L, "W": W, "T": T})
    return out


# ---------------------------------------------------------------------------- MPR (woodWOP) TERVEZET

def mpr_text(rec, holes, side):
    L, W, T = rec["meret"]["kesz"]["L"], rec["meret"]["kesz"]["W"], rec["meret"]["kesz"]["T"]
    ops = []
    for i, h in enumerate(holes, 1):
        tag = f'{h["funkcio"]} {h["kapcsolat"]} TERVEZET'
        if side in ("A", "B") and h["oldal"] == side:
            x = h["u"]
            y = h["v"] if side == "A" else W - h["v"]
            ops.append(("<102 \\BohrVert\\", {"XA": x, "YA": y, "BM": "SS", "TI": h["melyseg"], "DU": h["d"], "AN": 1, "MI": 0,
                                              "AB": 0, "F_": "STANDARD", "EN": 1, "KM": tag}))
        elif side == "A" and h["oldal"] == "EL":
            e = h["el"]
            xa, ya, bm = {"E4": (0, h["v"], "XP"), "E2": (L, h["v"], "XM"), "E1": (h["u"], 0, "YP"), "E3": (h["u"], W, "YM")}[e]
            ops.append(("<103 \\BohrHoriz\\", {"XA": xa, "YA": ya, "ZA": h["n"], "BM": bm, "TI": h["melyseg"], "DU": h["d"], "AN": 1,
                                               "MI": 0, "AB": 0, "F_": "STANDARD", "EN": 1, "KM": tag}))
    if not ops:
        return None, 0
    fmt = lambda v: (f"{v:.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else str(v))
    lines = ["[H", 'VERSION="4.0"', 'HP="1"', 'IN="0"', "", "[001", f'L="{fmt(float(L))}"', f'B="{fmt(float(W))}"', f'D="{fmt(float(T))}"', "",
             "<100 \\WerkStck\\", 'LA="L"', 'BR="B"', 'DI="D"', "",
             "<101 \\Kommentar\\",
             f'KM="{rec["cimke_id"]} {side}-OLDAL FELUL - {rec["modell_revizio"]} - TERVEZET, NEM GEPRE KESZ (minta-MPR es szerszamtabla hianyzik)"', ""]
    for head, kv in ops:
        lines.append(head)
        for k, v in kv.items():
            lines.append(f'{k}="{fmt(v) if not isinstance(v, str) else v}"')
        lines.append("")
    return "\r\n".join(lines) + "\r\n", len(ops)


def parse_mpr(text):
    ws = re.search(r'\[001\s+L="([\d.]+)"\s+B="([\d.]+)"\s+D="([\d.]+)"', text)
    L, W, T = map(float, ws.groups())
    ops, cur = [], None
    for line in text.splitlines():
        if line.startswith("<"):
            cur = {"_op": line.strip(), "_bad": []}
            ops.append(cur)
        elif cur is not None and "=" in line:
            k, v = line.split("=", 1)
            cur[k] = v.strip('"')
            allowed = MPR_KEYS_VERT if "BohrVert" in cur["_op"] else MPR_KEYS_HOR if "BohrHoriz" in cur["_op"] else ["KM", "LA", "BR", "DI"]
            if k not in allowed:
                cur["_bad"].append(k)
    return (L, W, T), ops


# ---------------------------------------------------------------------------- alkatrészlap (PDF + DXF)

def outline_uv(rec, src):
    p = src[rec["id"]]
    if p.get("outline") and p.get("axis") == "y" and rec["helyi_rendszer"]["A_normal"] == "+Y":
        return [tuple(map(float, q)) for q in p["outline"]], [[tuple(map(float, q)) for q in c["outline"]] for c in p.get("cutouts", [])]
    L, W = rec["meret"]["kesz"]["L"], rec["meret"]["kesz"]["W"]
    return [(0, 0), (L, 0), (L, W), (0, W)], []


def draw_sheet(pdf, rec, holes, src, mpr_names, blokk=()):
    import matplotlib.pyplot as plt
    from matplotlib.patches import Circle, Polygon as MPoly, Rectangle
    L, W, T = rec["meret"]["kesz"]["L"], rec["meret"]["kesz"]["W"], rec["meret"]["kesz"]["T"]
    fig = plt.figure(figsize=(11.69, 8.27))
    ax = fig.add_axes([0.05, 0.30, 0.62, 0.62])
    outl, cuts = outline_uv(rec, src)
    ax.add_patch(MPoly(outl, closed=True, fc="#f3efe8", ec="k", lw=1.2))
    for c in cuts:
        ax.add_patch(MPoly(c, closed=True, fc="white", ec="k", lw=0.9, ls="--"))
    # élzárás (vastag vonal a befoglaló oldalán), él-jel
    E = rec["elek"]
    seg = {"E1": ((0, 0), (L, 0)), "E2": ((L, 0), (L, W)), "E3": ((0, W), (L, W)), "E4": ((0, 0), (0, W))}
    col = {"V": "#c0392b", "F": "#8e44ad", "L": "#d35400", "K": "#16a085", "0": None}
    off = max(L, W) * 0.035
    for e, (a, b) in seg.items():
        j = E[e]["jel"]
        if col.get(j):
            ax.plot([a[0], b[0]], [a[1], b[1]], color=col[j], lw=4, solid_capstyle="butt", zorder=3)
        mx, my = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
        dx, dy = {"E1": (0, -off), "E2": (off, 0), "E3": (0, off), "E4": (-off, 0)}[e]
        ax.text(mx + dx, my + dy, f"{e}: {j}", ha="center", va="center", fontsize=8, color=col.get(j) or "#555", zorder=4)
    # ütköző (E1+E4 sarok)
    s = max(L, W) * 0.03
    ax.add_patch(MPoly([(0, 0), (-s, -s * 0.4), (-s * 0.4, -s)], closed=True, fc="k"))
    ax.text(-s * 1.2, -s * 1.2, "0,0 ütköző E1+E4", fontsize=7, ha="right", va="top")
    # szálirány
    g = rec["szalirany"]["helyi_tengely"]
    if g:
        cx, cy = L * 0.5, W * 0.5
        if g == "U":
            ax.annotate("", xy=(cx + L * 0.15, cy), xytext=(cx - L * 0.15, cy), arrowprops=dict(arrowstyle="<->", color="#7f8c8d"))
        else:
            ax.annotate("", xy=(cx, cy + W * 0.15), xytext=(cx, cy - W * 0.15), arrowprops=dict(arrowstyle="<->", color="#7f8c8d"))
        ax.text(cx, cy, "erezet", fontsize=7, color="#7f8c8d", ha="center", va="bottom")
    # furatok
    for i, h in enumerate(holes, 1):
        if h["oldal"] == "A":
            ax.add_patch(Circle((h["u"], h["v"]), h["d"] / 2, fc="#2c3e50", ec="#2c3e50", zorder=5))
        elif h["oldal"] == "B":
            ax.add_patch(Circle((h["u"], h["v"]), h["d"] / 2, fc="none", ec="#2980b9", lw=1.0, ls="--", zorder=5))
        else:
            dd = {"E4": (1, 0), "E2": (-1, 0), "E1": (0, 1), "E3": (0, -1)}[h["el"]]
            x0, y0 = (0 if h["el"] == "E4" else L if h["el"] == "E2" else h["u"]), (0 if h["el"] == "E1" else W if h["el"] == "E3" else h["v"])
            ax.plot([x0, x0 + dd[0] * h["melyseg"]], [y0, y0 + dd[1] * h["melyseg"]], color="#27ae60", lw=3, zorder=5)
    # befoglaló méretek
    ax.annotate("", xy=(L, -off * 2.2), xytext=(0, -off * 2.2), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(L / 2, -off * 2.6, f"{L:g}", ha="center", va="top", fontsize=9)
    ax.annotate("", xy=(-off * 2.2, W), xytext=(-off * 2.2, 0), arrowprops=dict(arrowstyle="<->", lw=0.8))
    ax.text(-off * 2.6, W / 2, f"{W:g}", ha="right", va="center", fontsize=9, rotation=90)
    pad = max(L, W) * 0.12
    ax.set_xlim(-pad, L + pad)
    ax.set_ylim(-pad, W + pad)
    ax.set_aspect("equal")
    ax.axis("off")
    ax.set_title(f"A oldal felülről ({rec['lapoldal']['A']}) · U → jobbra, V → fel", fontsize=9, loc="left")
    # furattábla
    tx = fig.add_axes([0.69, 0.30, 0.29, 0.62])
    tx.axis("off")
    rows = [f"{'#':>3} {'old.':<4} {'X':>7} {'Y':>7} {'Ø':>4} {'mély':>4}  funkció"]
    for i, h in enumerate(holes, 1):
        lab = h["oldal"] if h["oldal"] != "EL" else h["el"]
        rows.append(f"{i:>3} {lab:<4} {h['u']:>7.1f} {h['v']:>7.1f} {h['d']:>4g} {h['melyseg']:>4g}  {h['funkcio'][:14]}")
    if len(rows) > 46:
        rows = rows[:45] + [f"… további {len(holes) - 44} furat a furatlistában (Excel)"]
    tx.text(0, 1, "\n".join(rows) if holes else "Nincs furat (vasalat-/kötésterv szerint).", family="DejaVu Sans Mono",
            fontsize=6.4, va="top")
    # pecsét
    st = fig.add_axes([0.05, 0.03, 0.93, 0.22])
    st.axis("off")
    st.add_patch(Rectangle((0, 0), 1, 1, fill=False, lw=1.2, transform=st.transAxes))
    m = rec["meret"]
    info = [
        (f"{rec['cimke_id']}  –  {rec['megnevezes']}", 13, "bold"),
        (f"Anyag: {rec['anyag']['kod']} {rec['anyag']['megnevezes']} · vastagság {T:g} mm · db {rec['db']} · modul {rec['modul']} · {rec['gyartasi_tipus']}", 9, "normal"),
        (f"Kész: {L:g} × {W:g} · élzárás előtt: {m['elzaras_elotti']['L']:g} × {m['elzaras_elotti']['W']:g} · nyers: "
         f"{m['nyers']['L'] if m['nyers']['L'] is not None else '–'} × {m['nyers']['W'] if m['nyers']['W'] is not None else '–'} mm", 9, "normal"),
        (f"Él-kód (E1-E2-E3-E4): {rec['el_kod']} · ütköző: E1+E4 · szálirány: {rec['szalirany']['helyi_tengely'] or 'nincs'} · "
         f"furat: A {sum(h['oldal'] == 'A' for h in holes)} / B {sum(h['oldal'] == 'B' for h in holes)} / él {sum(h['oldal'] == 'EL' for h in holes)}", 9, "normal"),
        (f"MPR: {', '.join(mpr_names) if mpr_names else '– (nincs furat)'} · Rev {rec['modell_revizio']} · {TODAY} · {STATUS}", 8.5, "normal"),
    ]
    if blokk:
        info.append((f"BLOKKOLÓ ({len(blokk)}): " + blokk[-1][:200], 8, "bold"))
        fig.text(0.36, 0.95, "BLOKKOLÓ – NEM GYÁRTHATÓ, DÖNTÉS KELL", color="#c0392b", fontsize=14, fontweight="bold", ha="center")
    if rec.get("nyitott"):
        info.append(("NYITOTT: " + " ".join(rec["nyitott"])[:220], 7.5, "normal"))
    for k, (t, fs, fw) in enumerate(info):
        st.text(0.01, 0.92 - k * 0.135, t, fontsize=fs, fontweight=fw, transform=st.transAxes, va="top")
    st.text(0.99, 0.06, "WSG A-902 · S16 médiafal · Claude (AI) / Kapusi Péter", fontsize=7, ha="right", transform=st.transAxes)
    pdf.savefig(fig)
    plt.close(fig)


def write_dxf(rec, holes, src, path):
    import ezdxf
    L, W = rec["meret"]["kesz"]["L"], rec["meret"]["kesz"]["W"]
    d = ezdxf.new("R2018", setup=True)
    d.header["$INSUNITS"] = 4
    for name, c in (("KONTUR", 7), ("KIVAGAS", 7), ("FURAT_A", 5), ("FURAT_B", 4), ("FURAT_EL", 3), ("ELZARAS", 1), ("SZOVEG", 7), ("MERET", 8)):
        d.layers.add(name, color=c)
    msp = d.modelspace()
    outl, cuts = outline_uv(rec, src)
    msp.add_lwpolyline(outl, close=True, dxfattribs={"layer": "KONTUR"})
    for c in cuts:
        msp.add_lwpolyline(c, close=True, dxfattribs={"layer": "KIVAGAS"})
    seg = {"E1": ((0, 0), (L, 0)), "E2": ((L, 0), (L, W)), "E3": ((0, W), (L, W)), "E4": ((0, 0), (0, W))}
    for e, (a, b) in seg.items():
        j = rec["elek"][e]["jel"]
        if j != "0":
            msp.add_line(a, b, dxfattribs={"layer": "ELZARAS", "lineweight": 70})
        msp.add_text(f"{e}:{j}", height=8, dxfattribs={"layer": "SZOVEG"}).set_placement(((a[0] + b[0]) / 2, (a[1] + b[1]) / 2))
    for h in holes:
        if h["oldal"] in ("A", "B"):
            msp.add_circle((h["u"], h["v"]), h["d"] / 2, dxfattribs={"layer": f"FURAT_{h['oldal']}"})
        else:
            dd = {"E4": (1, 0), "E2": (-1, 0), "E1": (0, 1), "E3": (0, -1)}[h["el"]]
            x0 = 0 if h["el"] == "E4" else L if h["el"] == "E2" else h["u"]
            y0 = 0 if h["el"] == "E1" else W if h["el"] == "E3" else h["v"]
            msp.add_line((x0, y0), (x0 + dd[0] * h["melyseg"], y0 + dd[1] * h["melyseg"]), dxfattribs={"layer": "FURAT_EL"})
    dim = msp.add_linear_dim(base=(0, -40), p1=(0, 0), p2=(L, 0), dimstyle="EZDXF", dxfattribs={"layer": "MERET"})
    dim.render()
    dim = msp.add_linear_dim(base=(-40, 0), p1=(0, 0), p2=(0, W), angle=90, dimstyle="EZDXF", dxfattribs={"layer": "MERET"})
    dim.render()
    msp.add_text(f"{rec['cimke_id']} {rec['anyag']['kod']} T{rec['meret']['kesz']['T']:g} el-kod {rec['el_kod']} - {STATUS}",
                 height=10, dxfattribs={"layer": "SZOVEG"}).set_placement((0, -80))
    d.saveas(path)


# ---------------------------------------------------------------------------- címkék

def write_labels(parts, holes_by, mpr_by, path, blokk_by=None):
    from reportlab.graphics.barcode import code128
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    pdfmetrics.registerFont(TTFont("DV", FONT))
    pdfmetrics.registerFont(TTFont("DVB", FONT_B))
    c = canvas.Canvas(str(path), pagesize=(100 * mm, 50 * mm))
    for r in parts:
        m = r["meret"]["kesz"]
        bc = code128.Code128(r["cimke_id"], barHeight=11 * mm, barWidth=0.33 * mm, humanReadable=False)
        bc.drawOn(c, 4 * mm, 34.5 * mm)
        c.setFont("DVB", 10)
        c.drawString(4 * mm, 30 * mm, r["cimke_id"])
        c.setFont("DV", 7)
        c.drawString(4 * mm, 26.5 * mm, r["megnevezes"][:48])
        c.drawString(4 * mm, 23 * mm, f"{r['anyag']['kod']} {r['anyag']['megnevezes'][:38]}")
        c.setFont("DVB", 9)
        c.drawString(4 * mm, 18.5 * mm, f"{m['L']:g} × {m['W']:g} × {m['T']:g} mm")
        c.setFont("DV", 7)
        h = holes_by.get(r["id"], [])
        c.drawString(4 * mm, 14.5 * mm, f"él E1-E4: {r['el_kod']}   szál: {r['szalirany']['helyi_tengely'] or '–'}   A: {r['lapoldal']['A'].split(' ')[0]}")
        c.drawString(4 * mm, 11 * mm, f"furat A/B/él: {sum(x['oldal'] == 'A' for x in h)}/{sum(x['oldal'] == 'B' for x in h)}/{sum(x['oldal'] == 'EL' for x in h)}   modul: {r['modul']}")
        c.drawString(4 * mm, 7.5 * mm, f"MPR: {', '.join(mpr_by.get(r['id'], [])) or '–'}"[:70])
        if blokk_by and blokk_by.get(r["id"]):
            c.setFillColorRGB(0.75, 0.1, 0.1)
            c.setFont("DVB", 8)
            c.drawString(60 * mm, 30 * mm, "BLOKKOLÓ – NE GYÁRTSD")
            c.setFillColorRGB(0, 0, 0)
        c.setFont("DV", 5.5)
        c.drawString(4 * mm, 3.5 * mm, f"WSG A-902 S16 · {r['modell_revizio']} · {TODAY} · TERVI – NEM GYÁRTÁSI")
        # tájolás: kis téglalap az E1/E4 ütköző sarokkal
        x0, y0, w, hh = 80 * mm, 14 * mm, 15 * mm, 10 * mm
        c.rect(x0, y0, w, hh)
        c.setFont("DV", 5)
        c.drawString(x0 + w / 2 - 2 * mm, y0 - 2.5 * mm, "E1")
        c.drawString(x0 - 3 * mm, y0 + hh / 2, "E4")
        c.circle(x0, y0, 0.9 * mm, fill=1)
        c.showPage()
    c.save()


# ---------------------------------------------------------------------------- Excel

def write_excel(doc, parts, holes_by, mpr_by, vas, path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    F = "Arial"
    hfont, bfont = Font(name=F, bold=True, color="FFFFFF", size=10), Font(name=F, size=10)
    hfill = PatternFill("solid", fgColor="3B3530")
    blue = Font(name=F, size=10, color="0000FF")
    thin = Side(style="thin", color="BFBFBF")
    wb = Workbook()

    def header(ws, cols, widths):
        ws.append(cols)
        for i, (c, w) in enumerate(zip(cols, widths), 1):
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = hfont, hfill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 32

    # 1. Szabásjegyzék
    ws = wb.active
    ws.title = "Szabásjegyzék"
    cols = ["Címke-ID (MPR alap)", "Megnevezés", "Modul", "Gyártási típus", "Anyag", "Anyag megnevezés", "Vastagság mm",
            "Kész hossz L", "Kész szél. W", "Élzárás előtt L", "Élzárás előtt W", "Nyers L", "Nyers W", "Db", "Szálirány",
            "A oldal", "E1", "E2", "E3", "E4", "Élanyag (élzárt élek)", "Ütköző élek", "Furat A", "Furat B", "Furat él",
            "MPR fájlok", "Felület m²", "Élzáró fm", "Nyitott / megjegyzés"]
    header(ws, cols, [20, 30, 13, 15, 7, 34, 9, 9, 9, 10, 10, 9, 9, 5, 9, 10, 5, 5, 5, 5, 30, 9, 7, 7, 7, 26, 9, 9, 50])
    for r in parts:
        m = r["meret"]
        h = holes_by.get(r["id"], [])
        el = r["elek"]
        anyag = sorted({el[e]["elanyag"] for e in el if el[e]["jel"] in ("V", "F", "L", "K")})
        ws.append([r["cimke_id"], r["megnevezes"], r["modul"], r["gyartasi_tipus"], r["anyag"]["kod"], r["anyag"]["megnevezes"],
                   m["kesz"]["T"], m["kesz"]["L"], m["kesz"]["W"], m["elzaras_elotti"]["L"], m["elzaras_elotti"]["W"],
                   m["nyers"]["L"], m["nyers"]["W"], r["db"], r["szalirany"]["helyi_tengely"] or "–", r["lapoldal"]["A"],
                   el["E1"]["jel"], el["E2"]["jel"], el["E3"]["jel"], el["E4"]["jel"], "; ".join(anyag) or "–", "E1+E4",
                   sum(x["oldal"] == "A" for x in h), sum(x["oldal"] == "B" for x in h), sum(x["oldal"] == "EL" for x in h),
                   ", ".join(mpr_by.get(r["id"], [])) or "–", None, None, " ".join(r.get("nyitott", []))])
        i = ws.max_row
        ws[f"AA{i}"] = f"=H{i}*I{i}/1000000*N{i}"
        # élzáró folyóméter: E1/E3 a hosszél (L), E2/E4 a szélességél (W); csak a V (ABS) éleket számoljuk
        ws[f"AB{i}"] = f'=(IF(Q{i}="V",H{i},0)+IF(S{i}="V",H{i},0)+IF(R{i}="V",I{i},0)+IF(T{i}="V",I{i},0))/1000*N{i}'
    n = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.font = bfont
            c.border = Border(bottom=thin)
    for col in ("AA", "AB"):
        for i in range(2, n + 1):
            ws[f"{col}{i}"].number_format = "0.000"
    ws.add_table(Table(displayName="Szabasjegyzek", ref=f"A1:AC{n}", tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True)))
    ws[f"Z{n + 2}"] = "Összesen:"
    ws[f"AA{n + 2}"] = f"=SUM(AA2:AA{n})"
    ws[f"AB{n + 2}"] = f"=SUM(AB2:AB{n})"
    ws["AB1"].comment = Comment("Csak a V (ABS/élfólia) éleket számolja; F festett, L lakkozott, K kárpitozott él nem élzáró-anyag.", "Claude")
    ws["L1"].comment = Comment("Nyers = kész − élvastagság + előmarás + ráhagyás, a gyartasi_szabalyok_v1.json szerint (becslés, műhelyi megerősítésre).", "Claude")

    # 2. Anyag-összesítő
    wa = wb.create_sheet("Anyag-összesítő")
    header(wa, ["Anyag", "Megnevezés", "Darab", "Felület m²", "Élzáró fm (V)", "Tábla m² (becslés)", "Hulladék % (becslés)", "Tábla db (becslés)"],
           [8, 44, 8, 12, 13, 16, 16, 16])
    mats = {}
    for r in parts:
        mats.setdefault(r["anyag"]["kod"], r["anyag"]["megnevezes"])
    for k, (kod, nev) in enumerate(sorted(mats.items()), start=2):
        wa.append([kod, nev])
        wa[f"C{k}"] = f'=SUMIFS(Szabásjegyzék!N:N,Szabásjegyzék!E:E,A{k})'
        wa[f"D{k}"] = f'=SUMIFS(Szabásjegyzék!AA:AA,Szabásjegyzék!E:E,A{k})'
        wa[f"E{k}"] = f'=SUMIFS(Szabásjegyzék!AB:AB,Szabásjegyzék!E:E,A{k})'
        wa[f"F{k}"] = 0 if kod in ("M05", "M12") else 5.796
        wa[f"G{k}"] = 0.2
        wa[f"H{k}"] = f"=IF(F{k}>0,ROUNDUP(D{k}*(1+G{k})/F{k},0),0)"
        for c in ("F", "G"):
            wa[f"{c}{k}"].font = blue
        wa[f"D{k}"].number_format = wa[f"E{k}"].number_format = "0.00"
        wa[f"G{k}"].number_format = "0%"
    last = wa.max_row
    wa[f"A{last + 2}"] = "Feltételezés: tábla 2800×2070 mm = 5,796 m², hulladék 20 % – becslés, termék nincs kiválasztva (kék = módosítható). M05 szövet és M12 fénytakaró nem táblaanyag (0 = nem számol)."
    wa[f"A{last + 2}"].font = Font(name=F, italic=True, size=9)

    # 3. Vasalat
    wv = wb.create_sheet("Vasalat")
    header(wv, ["Tétel", "Cikkszám / megnevezés", "Db", "Forrás", "Státusz"], [26, 60, 8, 40, 46])
    cnt = Counter(w["tipus"] for w in vas["vasalat"])
    names = {w["tipus"]: w["cikkszam"] for w in vas["vasalat"]}
    for t, dbn in sorted(cnt.items()):
        db = sum(w.get("db", 1) for w in vas["vasalat"] if w["tipus"] == t)
        wv.append([t, names[t], db, "f2b_vasalat.py – geometriából", "TERVI – típus/variáns műhelyi megerősítésre"])
    wv.append(["PANT talplemez", "Blum CLIP top talplemez (variáns nyitott)", cnt.get("PANT_CLIPTOP_110", 0), "f2b_vasalat.py", "TERVI"])
    for row in [("Blum MOVENTO 760H sínpár (NL300)", "760H3000S pár + T51.7601 kapocs (R14 drawers.json)", 6, "R14 drawers.json", "NYITOTT – gyári furatkép a sín kiválasztása után"),
                ("TIP-ON BLUMOTION egység", "T60L7140 / T60L7040 (R14 drawers.json)", 6, "R14 drawers.json", "NYITOTT"),
                ("Bal toronyajtó pánt", "szabványos CLIP top NEM építhető be (31 mm keretszár)", 6, "f2b_vasalat.py BLOKKOLÓ", "DÖNTÉS KELL: csapos/keskeny keretes pánt vagy keret szélesítése"),
                ("Korpusz él-él illesztések kötése", "Häfele kötő / csavar (etalon: 637.76.333/351)", None, "vasalat_jelentes.md", "NYITOTT – műhelyi döntés")]:
        wv.append(list(row))
    for row in wv.iter_rows(min_row=2):
        for c in row:
            c.font = bfont

    # 4. Furatok
    wf = wb.create_sheet("Furatok")
    header(wf, ["Címke-ID", "#", "Oldal / él", "X (U) mm", "Y (V) mm", "Z (N) mm", "Ø mm", "Mélység mm", "Funkció", "Kapcsolat", "MPR"],
           [20, 5, 9, 10, 10, 10, 7, 10, 16, 36, 24])
    for r in parts:
        for i, h in enumerate(holes_by.get(r["id"], []), 1):
            side = h["oldal"] if h["oldal"] != "EL" else h["el"]
            mpr = f"{r['cimke_id']}_{'B' if h['oldal'] == 'B' else 'A'}.mpr"
            wf.append([r["cimke_id"], i, side, h["u"], h["v"], h["n"], h["d"], h["melyseg"], h["funkcio"], h["kapcsolat"], mpr])
    for row in wf.iter_rows(min_row=2):
        for c in row:
            c.font = bfont

    # 5. Nyitott kérdések + jelmagyarázat
    wn = wb.create_sheet("Nyitott és jelmagyarázat")
    wn.column_dimensions["A"].width = 26
    wn.column_dimensions["B"].width = 110
    items = [("STÁTUSZ", STATUS + f" · {doc['modell_revizio']} · {TODAY}"),
             ("Koordináták", "X = helyi U (hossz), Y = helyi V (szélesség), az A oldal felől nézve; 0,0 = E1+E4 sarok (javasolt gépi ütköző)."),
             ("Él-kód", "E1-E2-E3-E4 sorrend (E1 alsó, E2 jobb, E3 felső, E4 bal az A oldal felől). V = élzárt (ABS), F = festett, L = lakkozott, K = kárpitozott, 0 = nincs."),
             ("Asztalos konvenció", "Az etalon él-kódjának (pl. V0VV) sorrendje és az 'A' jel jelentése még nincs megerősítve – a válasz után átállítható."),
             ("MPR", "woodWOP TERVEZET (VERSION 4.0 formátum, csak furatok). Gépre kész csak az asztalos minta-MPR-je, szerszámtáblája és próbadarab után."),
             ("BLOKKOLÓ", "Bal toronyajtó (T-DOOR-L): 31 mm-es keretszár, Ø35 pántkehely nem fér – pántmegoldás vagy keretszélesítés kell."),
             ("Nyitott", "Lapanyag/dekor termék és táblaméret; élzáró vastagság és előmarás; polctartó típus; fiók sín furatképe; LED-profilok; helyszíni felmérés.")]
    for k, v in items:
        wn.append([k, v])
    for row in wn.iter_rows():
        row[0].font = Font(name=F, bold=True, size=10)
        row[1].font = bfont
        row[1].alignment = Alignment(wrap_text=True)
    wb.save(path)


# ---------------------------------------------------------------------------- main + QA

def main():
    from matplotlib.backends.backend_pdf import PdfPages
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import font_manager
    font_manager.fontManager.addfont(FONT)
    matplotlib.rcParams["font.family"] = "DejaVu Sans"

    doc, rules, vas, src = load_all()
    holes_by = classify_holes(doc, rules, vas)
    blokk_by = defaultdict(list)
    for pr in vas["problemak"]:
        if pr["sulyossag"] == "BLOKKOLO":
            blokk_by[pr["alkatresz"]].append(pr["uzenet"])
    parts = [r for r in doc["alkatreszek"]]
    for r in parts:
        r["modell_revizio"] = doc["modell_revizio"]
    (HERE / "mpr_tervezet").mkdir(exist_ok=True)
    (HERE / "dxf_alkatreszek").mkdir(exist_ok=True)
    mpr_by, mpr_counts = {}, {}
    for r in parts:
        hs = holes_by.get(r["id"], [])
        for side in ("A", "B"):
            txt, nops = mpr_text(r, hs, side)
            if txt:
                fn = f"{r['cimke_id']}_{side}.mpr"
                (HERE / "mpr_tervezet" / fn).write_text(txt, encoding="cp1252", newline="")
                mpr_by.setdefault(r["id"], []).append(fn)
                mpr_counts[fn] = nops

    with PdfPages(HERE / "S16_R15_alkatreszlapok.pdf") as pdf:
        for r in parts:
            draw_sheet(pdf, r, holes_by.get(r["id"], []), src, mpr_by.get(r["id"], []), blokk_by.get(r["id"], []))
    for r in parts:
        write_dxf(r, holes_by.get(r["id"], []), src, HERE / "dxf_alkatreszek" / f"{r['cimke_id']}.dxf")
    write_labels(parts, holes_by, mpr_by, HERE / "S16_R15_cimkek.pdf", blokk_by)
    write_excel(doc, parts, holes_by, mpr_by, vas, HERE / "S16_R15_szabasjegyzek.xlsx")

    # ------------------------------------------------ visszaolvasó QA
    qa = []
    total_v = sum(len(v) for v in holes_by.values())
    qa.append(("Vasalatterv furatai", len(vas["furatok"]), "alkatrészhez rendelve", total_v, len(vas["furatok"]) == total_v))
    nmpr = 0
    bad = []
    for fn in sorted((HERE / "mpr_tervezet").glob("*.mpr")):
        (L, W, T), ops = parse_mpr(fn.read_text(encoding="cp1252"))
        for o in ops:
            if o["_bad"]:
                bad.append(f"{fn.name}: ismeretlen kulcs {o['_bad']}")
            if "Bohr" in o["_op"]:
                nmpr += 1
                x, y, ti, du = float(o["XA"]), float(o["YA"]), float(o["TI"]), float(o["DU"])
                if not (-1e-6 <= x <= L + 1e-6 and -1e-6 <= y <= W + 1e-6):
                    bad.append(f"{fn.name}: furat a darabon kívül X{x} Y{y}")
                if "BohrVert" in o["_op"] and ti >= T:
                    bad.append(f"{fn.name}: függőleges furat mélység {ti} ≥ vastagság {T}")
                if "BohrHoriz" in o["_op"]:
                    za = float(o["ZA"])
                    if not (du / 2 <= za <= T - du / 2):
                        bad.append(f"{fn.name}: vízszintes furat ZA {za} kilóg a lapvastagságból")
    qa.append(("MPR-ből visszaolvasott furat", total_v, "MPR műveletek", nmpr, nmpr == total_v))
    import openpyxl
    wb = openpyxl.load_workbook(HERE / "S16_R15_szabasjegyzek.xlsx")
    xl_rows = wb["Szabásjegyzék"].max_row - 1 - 0
    xl_parts = sum(1 for r in wb["Szabásjegyzék"].iter_rows(min_row=2, values_only=True) if r[0] and str(r[0]).startswith("S16_"))
    xl_holes = wb["Furatok"].max_row - 1
    qa.append(("Excel alkatrész-sor", len(parts), "Szabásjegyzék", xl_parts, xl_parts == len(parts)))
    qa.append(("Excel furat-sor", total_v, "Furatok lap", xl_holes, xl_holes == total_v))
    neg = [r["cimke_id"] for r in parts for k in ("L", "W", "T") if not (isinstance(r["meret"]["kesz"][k], (int, float)) and r["meret"]["kesz"][k] > 0)]
    qa.append(("Negatív/nulla méret", 0, "kész méretek", len(neg), not neg))
    from pypdf import PdfReader
    npages = len(PdfReader(str(HERE / "S16_R15_alkatreszlapok.pdf")).pages)
    nlab = len(PdfReader(str(HERE / "S16_R15_cimkek.pdf")).pages)
    qa.append(("Alkatrészlap (PDF oldal)", len(parts), "oldal", npages, npages == len(parts)))
    qa.append(("Címke (PDF oldal)", len(parts), "oldal", nlab, nlab == len(parts)))
    ndxf = len(list((HERE / "dxf_alkatreszek").glob("*.dxf")))
    qa.append(("Alkatrész-DXF", len(parts), "fájl", ndxf, ndxf == len(parts)))
    ok = all(q[4] for q in qa) and not bad
    L = [f"# S16 F3–F4 dokumentáció – visszaolvasó QA ({doc['modell_revizio']}, {TODAY})", "",
         f"**Eredmény: {'MINDEN EGYEZIK' if ok else 'ELTÉRÉS VAN'}**", "", "| Ellenőrzés | elvárt | forrás | kapott | OK |", "|---|---|---|---|---|"]
    L += [f"| {a} | {b} | {c} | {d} | {'✓' if e else '✗'} |" for a, b, c, d, e in qa]
    L += ["", f"MPR-fájlok: {len(mpr_counts)} db · műveletek: {nmpr} · tiltott/ismeretlen kulcs: {sum('ismeretlen' in b for b in bad)}", ""]
    if bad:
        L += ["## Hibák", ""] + [f"- {b}" for b in bad[:50]]
    L += ["", "Hatókör: ez a fájlok egymás közti egyezését bizonyítja (adatmodell ↔ vasalatterv ↔ MPR ↔ Excel ↔ PDF ↔ DXF), "
          "nem a gépi futtathatóságot – az a woodWOP-ban, a műhely szerszámtáblájával és próbadarabbal igazolható."]
    (HERE / "dokumentacio_qa.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    print("\n".join(L[:14]))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
