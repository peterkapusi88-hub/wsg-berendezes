from pathlib import Path
import json,hashlib,csv,shutil
from openpyxl import load_workbook
P=Path(__file__).resolve().parents[1];O=P/'CP3';D=json.loads((O/'data/design_CP3.json').read_text());old=json.loads((P/'render_handoff/manifest_R02.json').read_text());runs=json.loads((O/'blender/render_runs.json').read_text());photos=json.loads((O/'evidence/product_photo_manifest.json').read_text());checks=json.loads((O/'blender/checks.json').read_text());R=D['revision']
dest=O/'render_handoff';(dest/'prompts').mkdir(exist_ok=True)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
blend='CP3/blender/A902_Kyoto_Ink_CP3.blend';blendsha=sha(P/blend)
units={u['id']:u for u in D['units']}
wb=load_workbook(P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx',data_only=True,read_only=True)
live={r[2]:r for r in wb['Elemregiszter'].iter_rows(min_row=2,values_only=True)}
camera_review={'V12':'Original camera is too close to the new wardrobe fronts; use F01 for technical inspection. A revised aesthetic camera needs a new ID, not a silent camera move.','V13':'Original camera is too close to the new wardrobe fronts; use F02 for technical inspection. A revised aesthetic camera needs a new ID, not a silent camera move.'}
def attach(path,role):return dict(path=path,role=role,exists=(P/path).is_file(),sha256=sha(P/path) if (P/path).is_file() else None)
def unit_text(u):
 txt=f"Custom structure {u['id']}: overall {u['w']} W × {u['d']} D × {u['h']} H mm; 18 mm carcass, specified front finish. Preserve the closed-front divisions from its CP3 drawing. "
 for b in u['cp3_bays']:
  txt+=f"{b['id']}: clear width {b['clear_w']} mm, clear depth {b['clear_d']} mm; shelf elevations {b['shelves_mm']} mm; rail elevations {b['rails_mm']} mm; lift rails {b['lift_rails_mm']} mm; {b['drawers']} drawers, proposed net internal size {b['drawer_inner_mm']} mm; internal lights at {b['lights_mm']} mm. "
 return txt
views=[]
for v in old:
 rr=next(r for r in runs if r['id']==v['view_id']+'_CP3');v=dict(v);v.update(revision=R,blender_file=blend,blender_sha256=blendsha,camera=rr,primary_control=rr['image'])
 v['products']=[dict(prod) for prod in v['products'] if prod['element'] not in ['KI-048','KI-049']]
 for prod in v['products']:
  if prod['element'] in live:
   row=live[prod['element']];prod.update(name=row[7],brand=row[8],sku=row[9],product_url=row[10],dimensions=row[11],price_status=row[21])
 refs=[attach(rr['image'],'Primary exact-camera native Blender control; simplified product meshes, no final-product certification'),attach('CP3/drawings/P-CP3.pdf','Whole-apartment placement and internal-light overlay')]
 for a in v['attachments']:
  if a['path'].startswith('targets/'):
   primary='CP3/sources/'+Path(a['path']).name
   refs.append(attach(primary if (P/primary).is_file() else a['path'],'Aesthetic target only, never geometric or dimensional evidence'))
 for uid in v['custom_structures']:
  if uid in units:refs.append(attach(units[uid]['drawings_cp3'],'CP3 custom structure: page1 dimensions/interior; page2 operation and clear compartments'))
 for prod in v['products']:
  ph=next((p for p in photos if p['element']==prod['element'] and 'file' in p),None)
  if ph:refs.append(attach(ph['file'],'Official exact-product-page photograph; check pictured variant before final use'));prod['photo']=ph['file'];prod['photo_sha256']=ph['sha256'];prod['asset_status']='PHOTO_PACKAGED; manufacturer3D not provided'
  else:prod['asset_status']='NO_EXACT_ASSET; stop final-product rendering if this item is visible'
 v['attachments']=[dict(order=i+1,**a) for i,a in enumerate(refs)]
 v['statuses']=dict(geometry='conditional; see executed CP3 checks',product='partial photographs; simplified3D',function='open gates, including kitchen and beds',cost='linked to172-row CP3 workbook; budget exceeded',professional='survey and approvals pending')
 v['release']='CONDITIONAL_DIAGNOSTIC_ONLY';v['open_conditions']=['Kitchen840 mm working aisle and open-dishwasher240 mm passage unresolved','Main05.2 desk/door and05.3 bed/chair decisions remain open','Master NORDLI drawer/nightstand interference unresolved','Several exact product meshes/SKUs and professional approvals missing']
 if v['view_id'] in camera_review:
  v['open_conditions'].append(camera_review[v['view_id']]);v['release']='BLOCKED_CAMERA_FRAMING';v['statuses']['camera']=camera_review[v['view_id']]
 v['output_file']=f"A902_KyotoInk_{v['view_id']}_{R}_CONDITIONAL_v1.png";views.append(v)
for vid,cid,room,title in [('V27','CP3_Cam_AT8','AT8-01','Eighth-floor reserve pantry and freezer'),('V28','CP3_Cam_T026','T-0-26','Garage-level seasonal and luggage storage')]:
 rr=next(r for r in runs if r['camera_id']==cid);st=next(s for s in json.loads((O/'data/storage_CP3.json').read_text()) if s['id']==room)
 refs=[attach(rr['image'],'Native furnished room study; planning envelope, not field-verified room'),attach('CP3/drawings/ST_'+room+'.pdf','Storage plan with uncertainty and furniture IDs')]
 for ph in photos:
  if ph['element'] in (['KI-CP3-S01'] if room=='AT8-01' else ['KI-CP3-S02']) and 'file' in ph:refs.append(attach(ph['file'],'Exact product-page shelving photograph'))
 views.append(dict(view_id=vid,camera_id=cid,room=room,title=title,revision=R,blender_file=blend,blender_sha256=blendsha,camera=rr,primary_control=rr['image'],attachments=[dict(order=i+1,**a) for i,a in enumerate(refs)],products=st['items'],custom_structures=[],light_scene='Neutral diagnostic illumination; no verified installed storage luminaire',statuses=dict(geometry='planning envelope, finished dimensions unverified',product='dimensioned shelving/freezer envelope; exact freezer EAN pending',function='door and ventilation survey required',cost='CP3 includes shelves, logistics, boxes; freezer included in pictured main case',professional='building and electrical checks pending'),release='CONDITIONAL_DIAGNOSTIC_ONLY',open_conditions=[st['dimension_status'],st['door'],'Do not add water outlets or assume a garage electrical socket'],output_file=f'A902_KyotoInk_{vid}_{R}_CONDITIONAL_v1.png'))
def prompt(v):
 lines=[f"A902 KYOTO INK — {v['view_id']}: {v['title']}.",f"Create only a CONDITIONAL DESIGN STUDY for room {v['room']}, revision {R}. This is not an approved construction photograph.",f"Use Blender file {blend}, SHA256 {blendsha}. Camera: {v['camera_id']}. Copy its exact camera transform from the attached control; do not reuse another perspective under this camera ID.",f"Output filename: {v['output_file']}. Lighting scene: {v['light_scene']}.","Attachment order and role:"]
 for a in v['attachments']:lines.append(f"{a['order']}. {a['path']} — {a['role']}. SHA256 {a['sha256']}.")
 lines+=['Visible product and custom-structure identity:']
 for prod in v['products']:lines.append(json.dumps(prod,ensure_ascii=False))
 for uid in v['custom_structures']:
  if uid in units:lines.append(unit_text(units[uid]))
 lines+=['Materials: retain Kyoto Ink matte graphite lower cabinetry, chalk/warm-white upper fronts, taupe textiles, light existing floor, and the specified ink artwork. Do not change any surface merely to conceal an unresolved product choice. Use the target image for aesthetic direction and the CP3 drawings for geometry. Match panel joints, service covers, and open/closed states exactly.',
 'Internal wardrobe light is separately switched functional lighting. Decorative/general lighting remains Hue. Do not invent visible glow without a modelled fixture/profile; CRI and colour temperature remain subject to manufacturer data and a sample.',
 'Geometry limits and release conditions: '+ '; '.join(v['open_conditions']),
 'Keep thresholds, window opening types, wall positions, furniture footprints, storage contents and the terrace boundary. Do not stretch a kitchen aisle or reduce a purchased appliance to make it fit. A visible simplified proxy cannot be made into a supposedly verified exact product by generative guessing.',
 'Before delivery compare the result to the native control and the exact product references, record any discrepancy, and save both the output and its audit. Stop on geometry drift, substituted products, changed front divisions or invented wiring. Never label the result final while any required release condition remains open.']
 return '\n\n'.join(lines)
for v in views:
 file=dest/'prompts'/f"{v['view_id']}_{v['camera_id']}.txt";file.write_text(prompt(v));v['prompt_path']=str(file.relative_to(P));v['prompt_sha256']=sha(file)
functional=[]
for r in runs:
 if not r['id'].startswith('F'):continue
 uid={'F01':'WN-01','F02':'WS-01','F03':'W52-01','F04':'W53-01','F05':'K-06','F06':'B62-01','F07':'ENT-01'}[r['id'][:3]]
 refs=[attach(r['image'],'Native functional state; use same camera for the closed/open/cutaway series'),attach(units[uid]['drawings_cp3'],'Pages1–2: dimensions and functional details')]
 txt=f"A902 Kyoto Ink {r['id']}, structure {uid}, {R}. Reproduce the exact native technical control state: {r['state']}. Use camera {r['camera_id']} with transform {json.dumps(r)}. Blender SHA256 {blendsha}.\n\n"+unit_text(units[uid])+"\n\nKeep this a labelled conditional technical visualization. A cutaway deliberately removes fronts; never present it as an ordinary open-door photograph. The lift shows an indicative end position, not certified linkage motion. Preserve dark clothing and representative boxes for shadow inspection. Do not add unmodelled shelves or hide a drawer conflict. Verify against both drawing pages and stop on identity/geometry drift.\n\nAttachments:\n"+'\n'.join(f'{i+1}. {a["path"]}, {a["role"]}, SHA256 {a["sha256"]}' for i,a in enumerate(refs))
 file=dest/'prompts'/(r['id']+'.txt');file.write_text(txt);functional.append(dict(**r,unit=uid,attachments=refs,prompt_path=str(file.relative_to(P)),release='CONDITIONAL_TECHNICAL_ONLY'))
(dest/'manifest_CP3.json').write_text(json.dumps(dict(revision=R,blender_sha256=blendsha,views=views,functional_views=functional),ensure_ascii=False,indent=2))
start='''You are the render operator for A902 Kyoto Ink CP3-R03. Continue the existing project; do not invent a new style. Read CP3/render_handoff/manifest_CP3.json and CP3/blender/checks.json before generating anything. The attached native Blender controls are the geometric authority; aesthetic targets are not dimensional evidence.

First perform ONE technical test using F01_closed and its two-page WN-01 drawing. Verify the exact attachment order and SHA256 hashes. Check image resolution, camera framing, panel divisions and product identity. Save the test plus an audit before any further image. A failed test stops the batch.

The package has V01–V28 room views and F01–F07 closed/open/cutaway functional views. Use each complete standalone English prompt. Never rename another perspective to an existing camera ID. Do not render optional room alternatives as approved without an explicit selection recorded in the workbook, drawings, model and manifest.

Every current view is conditional. Several final product meshes, field dimensions, opening clearances and technical approvals are unresolved. A request for final construction imagery must stop at those gates. Conditional studies may proceed only with clear labels and unchanged known limitations. Preserve actual thresholds and opening types. Never repair geometry by generative distortion, replace a cheaper item with a prettier one, or imply a fixture is Hue-compatible without evidence.

For each authorized image: load its exact control and references; compare output to control; check product variants, materials, artwork and storage interiors; save the output under its specified name; record hashes, status and discrepancies. Stop immediately if the result moves objects, changes front divisions, substitutes products or conceals a documented conflict. Do not order anything, contact suppliers or run mass generation automatically.
'''
(dest/'START_RENDER_OPERATOR_EN.txt').write_text(start)
brochure='''A902 KYOTO INK - KÉSŐBBI BROSÚRA BRIEF
Most nem kész prospektus. A jóváhagyott CP3 tervből épülő, magyar nyelvű tulajdonosi történet.
1. Miért érdekes? V02 nappali: tusgrafika, grafit/kréta ritmus, valóban rejtett tárolás.
2. Mi történik a frontok mögött? F01/F02 zárt és nyitott gardrób; nettó kapacitás és liftek.
3. Egy nap a lakásban: konyhai sorrend, munkavégzés, gyerek és kutya napi holmijai.
4. Két külön tároló: V27 tartalék+fagyasztó, V28 szezonális holmi; nem a napi gardrób helyett.
5. Ár és idő: aktuális Excelből, webshopár/ajánlat/becslés megkülönböztetve. Tartalék külön.
6. Hogyan induljon? Közös felmérés, mintacsomag, azonos revíziójú ajánlatok és legfeljebb3 döntés.
Csak ténylegesen kiválasztott termék és elfogadott alternatíva szerepelhet végleges képen. A conditional feliratot nem szabad elhagyni, amíg a kapuk nyitottak. Nincs bizonyítatlan95%-os egyezés vagy garantált ingatlanérték-növekedés.
Folytatóprompt: Olvasd vissza a legfrissebb CP3 manifestet és élő Excelt. Ellenőrizd a változások és jóváhagyások állapotát. A fenti hat jelenetből készíts tulajdonosi brosúrát kizárólag ellenőrzött képekkel és aktuális költséggel; minden megmaradt eltérést közérthetően jelezz.
'''
(O/'specifications/BROSURA_KESOBBI_BRIEF.txt').write_text(brochure)
print('28 teljes nézetprompt +21 funkcionális prompt + indítóprompt és manifest.')
