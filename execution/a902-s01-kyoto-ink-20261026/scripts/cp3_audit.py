from pathlib import Path
import json,hashlib,re
from collections import Counter
from openpyxl import load_workbook
P=Path(__file__).resolve().parents[1];O=P/'CP3';B=json.loads((P/'data/workbook_readback_before_CP3.json').read_text());patch=json.loads((O/'data/cost_patch.json').read_text());w=load_workbook(P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx',data_only=False);cached=load_workbook(P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx',data_only=True)
def col(i):
 s=''
 while i:i,n=divmod(i-1,26);s=chr(65+n)+s
 return s
changes=[];unexpected=[]
for name,rows in B.items():
 s=w[name]
 for ri,row in enumerate(rows,1):
  for ci,old in enumerate(row,1):
   new=s.cell(ri,ci).value
   if (old or None)==(new or None):continue
   a=col(ci)+str(ri);reason=None
   if name=='Elemregiszter' and ri>1 and col(ci) in patch['patches'].get(row[2],{}):reason='CP3 targeted patch'
   if isinstance(old,str) and old.startswith('=') and 'Elemregiszter' in old and re.search(r'\$?160\b',old) and new==re.sub(r'(\$?[A-Z]+\$?)160\b',r'\g<1>173',old):reason='sum range extended'
   if name=='Összesítő' and a in ['C7','C8','B17','C15','B15']:reason='CP3 summary'
   change=dict(sheet=name,cell=a,old=old,new=new,reason=reason);changes.append(change)
   if not reason:unexpected.append(change)
manifest=json.loads((O/'render_handoff/manifest_CP3.json').read_text());refs=[]
for view in manifest['views']+manifest['functional_views']:
 for a in view.get('attachments',[]):
  p=P/a['path'];ok=p.is_file() and hashlib.sha256(p.read_bytes()).hexdigest()==a['sha256'];refs.append(dict(path=a['path'],ok=ok))
checks=json.loads((O/'blender/checks.json').read_text());report=dict(changed_existing_cells=len(changes),unexpected_changes=unexpected,all_changes=changes,rows=w['Elemregiszter'].max_row-1,sheets=len(w.sheetnames),cached_main=cached['Összesítő']['B7'].value,cached_economy=cached['Összesítő']['B8'].value,attachments=len(refs),bad_attachments=[x for x in refs if not x['ok']],blender_check_counts=dict(Counter(c['result'] for c in checks)),camera_blocks=['V12','V13'],notes=['Product photographs visually checked as intended product family and colour; photographs do not prove size variant.','V12/V13 camera too close: explicitly blocked, F01/F02 functional cameras supplied.','All other native controls inspected as simplified geometry, not exact finished product renders.'])
(O/'qa/final_audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str));print(json.dumps({k:v for k,v in report.items() if k not in ['all_changes']},ensure_ascii=False,default=str))
