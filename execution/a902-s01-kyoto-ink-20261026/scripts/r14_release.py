from pathlib import Path
import json,re,hashlib,zipfile,shutil,xml.etree.ElementTree as ET,sqlite3,subprocess,sys,tempfile,struct
import openpyxl,ezdxf
from pypdf import PdfReader
P=Path(__file__).resolve().parents[1];O=P/'R14_muhely'
parts=json.loads((O/'data/parts.json').read_text());holes=json.loads((O/'data/holes.json').read_text());programs=json.loads((O/'data/programs.json').read_text())
shutil.copy2('/tmp/ki14_verify_final.log',O/'audit/autocad_readback.log');shutil.copy2('/tmp/ki14_verify.scr',O/'audit/autocad_readback.scr')
shutil.copy2('/tmp/ki14_build_final.log',O/'audit/autocad_build.log');shutil.copy2('/tmp/ki14_views_final.log',O/'audit/autocad_views.log')
log=(O/'audit/autocad_readback.log').read_text();volume=float(re.search(r'Volume:\s+([0-9.]+)',log).group(1));expected=json.loads((O/'audit/design_checks.json').read_text())['expected_net_volume_mm3']
assert '(KI_R14_SOLIDS 75)' in log and 'Total errors found 0 fixed 0' in log and abs(volume-expected)<.001
assert all(p['id'] in log for p in parts)
assert '(KI_R14_CUTS 579)' in (O/'audit/autocad_build.log').read_text()
(O/'audit/autocad_result.json').write_text(json.dumps(dict(engine='AutoCAD2027 AcCoreConsole',solids=75,subtracted_bores=579,volume_mm3=volume,analytic_volume_mm3=expected,volume_match=True,database_errors=0,not_certified=['field measurements','wall/hardware load capacity','hinge kinetics','machine/tool/fixture collision','woodWOP5 acceptance']),indent=2))
# Add print metadata unsupported by the authoring API; do not change workbook cells.
f=O/'TV_R14_alkatreszek_furatok_cimkek.xlsx';ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda s:'{'+ns+'}'+s;ET.register_namespace('',ns)
with zipfile.ZipFile(f) as z:items={n:z.read(n) for n in z.namelist()}
s=ET.fromstring(items['xl/worksheets/sheet5.xml'])
for n in ['pageMargins','pageSetup','rowBreaks']:
 e=s.find(q(n))
 if e is not None:s.remove(e)
ET.SubElement(s,q('pageMargins'),{'left':str(5/25.4),'right':str(5/25.4),'top':str(23.5/25.4),'bottom':str(23.5/25.4),'header':'0','footer':'0'})
ET.SubElement(s,q('pageSetup'),{'paperSize':'9','orientation':'portrait','scale':'100'})
b=ET.SubElement(s,q('rowBreaks'),{'count':'7','manualBreakCount':'7'})
for r in range(50,351,50):ET.SubElement(b,q('brk'),{'id':str(r),'min':'0','max':'16383','man':'1'})
items['xl/worksheets/sheet5.xml']=ET.tostring(s,encoding='utf-8',xml_declaration=True)
w=ET.fromstring(items['xl/workbook.xml']);dn=w.find(q('definedNames'))
if dn is None:dn=ET.Element(q('definedNames'));w.insert(list(w).index(w.find(q('sheets')))+1,dn)
for el in list(dn):
 if el.get('name')=='_xlnm.Print_Area' and el.get('localSheetId')=='4':dn.remove(el)
ET.SubElement(dn,q('definedName'),{'name':'_xlnm.Print_Area','localSheetId':'4'}).text="'Cimkek'!$A$1:$I$380"
items['xl/workbook.xml']=ET.tostring(w,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(f,'w',zipfile.ZIP_DEFLATED) as z:
 for n,v in items.items():z.writestr(n,v)
a=openpyxl.load_workbook(f,data_only=False);v=openpyxl.load_workbook(f,data_only=True)
errs=[[sh.title,cell.coordinate,cell.value] for sh in v for row in sh for cell in row if cell.data_type=='e'];assert not errs,errs
assert [v['Attekintes'][f'B{i}'].value for i in [2,3,4,5]]==[75,579,69,579]
assert abs(v['Koltsegnovekmeny']['I16'].value-126945.6)<.01
old=json.loads((O/'audit/R13_excel_snapshot.json').read_text());manual={r[0]:r for r in old['sheets']['Darabok'][5:] if r[0]}
for i,p in enumerate(parts,6):
 assert [a['Darabok'].cell(i,j).value for j in [1,3,4,5]]==[p['id'],p['x'],p['y'],p['t']]
 if p['id'] in manual:
  for j in list(range(7,13))+[16,17]:assert a['Darabok'].cell(i,j).value==manual[p['id']][j-1],(p['id'],j)
assert hashlib.sha256((O/'sources/TV_R13_eredeti_munkafuzet.xlsx').read_bytes()).hexdigest()==old['sha256']
book=PdfReader(O/'rajzok/TV_R14_reszletes_rajzkonyv.pdf');alltext='\n'.join(p.extract_text() for p in book.pages)
assert all(p['id'] in alltext for p in parts)
assert len(book.pages)==87 and len(PdfReader(O/'cimkek/TV_R14_cimkek_A4.pdf').pages)==8
for pr in programs:
 data=(O/pr['path']).read_bytes();assert hashlib.sha256(data).hexdigest()==pr['sha256'];assert b'EN="0"' not in data and b'??="0"' not in data
for p in parts:
 test=subprocess.run([sys.executable,str(O/'integration/cimke_kereso.py'),'R14:'+p['id']],capture_output=True,text=True);assert test.returncode==0,(p['id'],test.stderr)
stale=subprocess.run([sys.executable,str(O/'integration/cimke_kereso.py'),'R13:'+parts[0]['id']],capture_output=True);assert stale.returncode!=0
# Negative test on isolated copies: never tamper with deliverables.
with tempfile.TemporaryDirectory(prefix='ki14-integrity-') as td:
 t=Path(td);(t/'integration').mkdir();shutil.copy2(O/'integration/cimke_kereso.py',t/'integration');shutil.copy2(O/'integration/alkatreszek_R14.sqlite',t/'integration')
 db=sqlite3.connect(t/'integration/alkatreszek_R14.sqlite');row=db.execute('SELECT dxf,programs FROM parts WHERE barcode=?',('R14:'+parts[0]['id'],)).fetchone();db.close()
 files=[row[0]]+json.loads(row[1])
 for name in files:(t/name).parent.mkdir(exist_ok=True,parents=True);shutil.copy2(O/name,t/name)
 target=t/files[-1];target.write_bytes(target.read_bytes()+b'\nCORRUPTED_TEST\n')
 result=subprocess.run([sys.executable,str(t/'integration/cimke_kereso.py'),'R14:'+parts[0]['id']],capture_output=True,text=True);assert result.returncode!=0 and 'Hash eltérés' in result.stderr
dxfs=list((O/'dxf_alkatreszek').glob('*.dxf'))+list((O/'rajzok').glob('D0*.dxf'))+[O/'rajzok/TV_R14_LED_szerviz_javaslat.dxf']
for f in dxfs:assert not ezdxf.readfile(f).audit().has_errors,f
(O/'audit/release_checks.json').write_text(json.dumps(dict(part_ids=75,holes=579,active_mpr_files=69,all_macros_enabled=True,own_macro_readback=True,woodwop_test=False,machine_test=False,workbook_formula_errors=errs,workbook_cached_counts=[75,579,69,579],R13_manual_fields_preserved=True,pdf_pages=87,label_pages=8,label_lookup_all75=True,obsolete_revision_rejected=True,corrupted_file_rejected=True,physical_label_scan_test=False,editable_DXF_audits=len(dxfs),electrical_status='proposal only',lower_fronts_3_4_status='conflict unresolved, S10 separate alternative'),indent=2))
stls=[]
for p in parts:
 f=Path('/tmp/ki14_stls')/(p['id']+'.stl');data=f.read_bytes();n=struct.unpack_from('<I',data,80)[0];vs=[]
 for i in range(n):
  vals=struct.unpack_from('<12fH',data,84+50*i);vs.extend([vals[j:j+3] for j in [3,6,9]])
 lo=[min(v[k] for v in vs) for k in range(3)];hi=[max(v[k] for v in vs) for k in range(3)];box=p['assembly_box_mm']
 shift=[-box[k]+.000001 if box[k]<0 else 0 for k in range(3)]
 assert all(abs(lo[k]-shift[k]-box[k])<.001 and abs(hi[k]-shift[k]-box[k]-box[k+3])<.001 for k in range(3)),p['id']
 stls.append({'part':p['id'],'triangles':n,'sha256':hashlib.sha256(data).hexdigest(),'raw_bbox_mm':[lo,hi],'AutoCAD_STLOUT_positive_octant_translation_reversed_mm':shift,'restored_bbox_mm':[[lo[k]-shift[k] for k in range(3)],[hi[k]-shift[k] for k in range(3)]]})
(O/'audit/blender_source_provenance.json').write_text(json.dumps({'source_DWG':'rajzok/TV_R14_zart.dwg','source_DWG_sha256':hashlib.sha256((O/'rajzok/TV_R14_zart.dwg').read_bytes()).hexdigest(),'exports':stls,'source':'native AutoCAD STLOUT; documented negative-axis translation reversed; all nominal bounds checked','translation_reference':'https://www.autodesk.com/support/technical/article/caas/sfdcarticles/sfdcarticles/STLOUT-cannot-export-model-to-negative-coordinates-from-AutoCAD.html','render':'r14_preview.py / r14_detail_preview.py; no generative image edits'},indent=2))
manifest=[]
for f in sorted(O.rglob('*')):
 if not f.is_file() or f.name=='file_manifest.json' or f.name.endswith('.inspect.ndjson') or f.suffix in ['.dwl','.dwl2','.blend1']:continue
 manifest.append({'path':str(f.relative_to(O)),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(O/'data/file_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
zpath=P/'A902_TV_R14_muhelyi_checkpoint.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
 for x in manifest:z.write(O/x['path'],'R14_muhely/'+x['path'])
 z.write(O/'data/file_manifest.json','R14_muhely/data/file_manifest.json')
 for f in list((P/'scripts').glob('r14_*.py'))+[P/'r14_workbook.mjs']:z.write(f,'rebuild/'+f.name)
 z.writestr('rebuild/OLVASD_EL.txt','A generátorok az execution munkamappában futnak, a meglévő scripts/r13_*.py és R13_muhely forrásokkal. A csomag kész DWG/PDF/XLSX/DXF/MPR/Blender fájljai újragenerálás nélkül használhatók. AutoCAD/Blender futtatás a helyi telepítést igényli. A rebuild nem önálló környezettelepítő.')
with zipfile.ZipFile(zpath) as z:assert z.testzip() is None
print(json.dumps({'files':len(manifest),'zip_bytes':zpath.stat().st_size,'CAD_volume_match':True,'XLSX_errors':0,'PDF_pages':87,'all75lookups':True,'tampered_file_rejected':True}))
