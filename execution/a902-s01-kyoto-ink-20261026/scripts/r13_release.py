from pathlib import Path
import json,re,hashlib,zipfile,shutil,xml.etree.ElementTree as ET,sqlite3,subprocess
import openpyxl
from pypdf import PdfReader
P=Path(__file__).resolve().parents[1];O=P/'R13_muhely'
parts=json.loads((O/'data/parts.json').read_text());holes=json.loads((O/'data/holes.json').read_text());programs=json.loads((O/'data/programs.json').read_text())
shutil.copy('/tmp/ki13_verify.log',O/'audit/autocad_readback.log');shutil.copy('/tmp/ki13_verify.scr',O/'audit/autocad_readback.scr')
log=(O/'audit/autocad_readback.log').read_text();volume=float(re.search(r'Volume:\s+([0-9.]+)',log).group(1));expected=json.loads((O/'audit/design_checks.json').read_text())['expected_net_volume_mm3']
assert '(KI_R13_SOLIDS 54)' in log and 'Total errors found 0 fixed 0' in log and abs(volume-expected)<.001
(O/'audit/autocad_result.json').write_text(json.dumps(dict(engine='AutoCAD2027 AcCoreConsole X.119.M.246',solids=54,subtracted_bores=293,volume_mm3=volume,analytic_volume_mm3=expected,volume_match=True,database_errors=0,not_certified=['site dimensions','fastener loads','hinge kinematics','machine-tool collision','woodWOP5 parsing']),indent=2))
# Only print metadata missing in artifact-tool; workbook data/formulas stay intact.
f=O/'TV_R13_alkatreszek_furatok_cimkek.xlsx';ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda s:'{'+ns+'}'+s;ET.register_namespace('',ns)
with zipfile.ZipFile(f) as z:items={n:z.read(n) for n in z.namelist()}
s=ET.fromstring(items['xl/worksheets/sheet5.xml'])
for n in ['pageMargins','pageSetup','rowBreaks']:
 e=s.find(q(n))
 if e is not None:s.remove(e)
ET.SubElement(s,q('pageMargins'),{'left':str(5/25.4),'right':str(5/25.4),'top':str(23.5/25.4),'bottom':str(23.5/25.4),'header':'0','footer':'0'})
ET.SubElement(s,q('pageSetup'),{'paperSize':'9','orientation':'portrait','scale':'100'})
b=ET.SubElement(s,q('rowBreaks'),{'count':'5','manualBreakCount':'5'})
for r in [50,100,150,200,250]:ET.SubElement(b,q('brk'),{'id':str(r),'min':'0','max':'16383','man':'1'})
items['xl/worksheets/sheet5.xml']=ET.tostring(s,encoding='utf-8',xml_declaration=True)
w=ET.fromstring(items['xl/workbook.xml']);dn=w.find(q('definedNames'))
if dn is None:
 dn=ET.Element(q('definedNames'));w.insert(list(w).index(w.find(q('sheets')))+1,dn)
for el in list(dn):
 if el.get('name')=='_xlnm.Print_Area' and el.get('localSheetId')=='4':dn.remove(el)
ET.SubElement(dn,q('definedName'),{'name':'_xlnm.Print_Area','localSheetId':'4'}).text="'Cimkek'!$A$1:$I$270"
items['xl/workbook.xml']=ET.tostring(w,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(f,'w',zipfile.ZIP_DEFLATED) as z:
 for n,v in items.items():z.writestr(n,v)
a=openpyxl.load_workbook(f,data_only=False);v=openpyxl.load_workbook(f,data_only=True)
errs=[]
for sh in v:
 for row in sh:
  for cell in row:
   if cell.data_type=='e':errs.append([sh.title,cell.coordinate,cell.value])
assert not errs,errs
assert [v['Attekintes'][f'B{i}'].value for i in [2,3,4,5]]==[54,293,34,293]
for i,p in enumerate(parts,6):
 assert [a['Darabok'].cell(i,j).value for j in [1,3,4,5]]==[p['id'],p['x'],p['y'],p['t']]
assert str(a['Cimkek'].page_setup.paperSize)=='9' and len(a['Cimkek'].row_breaks.brk)==5
book=PdfReader(O/'rajzok/TV_R13_reszletes_rajzkonyv.pdf');alltext='\n'.join(p.extract_text() for p in book.pages)
assert all(p['id'] in alltext for p in parts)
assert len(book.pages)==63 and len(PdfReader(O/'cimkek/TV_R13_cimkek_A4.pdf').pages)==6
for pr in programs:
 data=(O/pr['path']).read_bytes();assert hashlib.sha256(data).hexdigest()==pr['sha256'];assert b'EN="0"' not in data and b'??="0"' not in data
db=sqlite3.connect(O/'integration/alkatreszek_R13.sqlite')
for p in parts:
 row=db.execute('SELECT dxf,programs FROM parts WHERE barcode=?',('R13:'+p['id'],)).fetchone();assert row
 for name in [row[0]]+json.loads(row[1]):assert (O/name).is_file()
assert not db.execute('SELECT * FROM parts WHERE barcode=?',('R12:KI12-TV01S-01',)).fetchone();db.close()
(O/'audit/release_checks.json').write_text(json.dumps(dict(part_ids=54,holes=293,active_mpr_files=34,all_macros_enabled=True,own_macro_readback=True,woodwop_test=False,machine_test=False,workbook_formula_errors=errs,workbook_cached_counts=[54,293,34,293],pdf_pages=63,label_pages=6,label_lookup_all54=True,obsolete_revision_rejected=True,physical_label_scan_test=False),indent=2))
manifest=[]
for f in sorted(O.rglob('*')):
 if not f.is_file() or f.name=='file_manifest.json' or f.name.endswith('.inspect.ndjson') or f.suffix in ['.dwl','.dwl2']:continue
 manifest.append({'path':str(f.relative_to(O)),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(O/'data/file_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
zpath=P/'A902_TV_R13_muhelyi_checkpoint.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
 for x in manifest:z.write(O/x['path'],'R13_muhely/'+x['path'])
 z.write(O/'data/file_manifest.json','R13_muhely/data/file_manifest.json')
 for f in list((P/'scripts').glob('r13_*.py'))+[P/'r13_workbook.mjs']:z.write(f,'rebuild/'+f.name)
with zipfile.ZipFile(zpath) as z:assert z.testzip() is None
print(json.dumps({'files':len(manifest),'zip_bytes':zpath.stat().st_size,'CAD_volume_match':True,'XLSX_errors':0,'PDF_pages':63}))
