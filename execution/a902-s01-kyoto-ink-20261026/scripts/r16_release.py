from pathlib import Path
import json,re,math,hashlib,zipfile,shutil,subprocess,sys,sqlite3
import openpyxl,ezdxf
from pypdf import PdfReader
P=Path(__file__).resolve().parents[1];O=P/'R16_feluletek_fiokok'
def r(n):return json.loads((O/n).read_text())
parts=r('data/parts.json');holes=r('data/holes.json');programs=r('data/programs.json');check=r('audit/design_checks.json')
log=(O/'audit/autocad_build.log').read_text();vol=float(re.search(r'Volume:\s+([0-9.]+)',log).group(1));assert abs(vol-check['expected_net_volume_mm3'])<.001
assert '(KI_R16_CUTS 956)' in log and 'Total errors found 0 fixed 0' in log
log=(O/'audit/autocad_saved_readback.log').read_text();assert '(R16_SOLIDS 215)' in log and '(R16_HATCHES 145)' in log and 'Total errors found 0 fixed 0' in log
assert len(parts)==106 and len(holes)==956 and len({h['id'] for h in holes})==956
assert sum(p['holes'] for p in programs)==956
for pr in programs:assert hashlib.sha256((O/pr['path']).read_bytes()).hexdigest()==pr['sha256']
for p in parts:
 t=subprocess.run([sys.executable,str(O/'integration/cimke_kereso.py'),'R16:'+p['id']],capture_output=True,text=True);assert t.returncode==0,(p['id'],t.stderr)
t=subprocess.run([sys.executable,str(O/'integration/cimke_kereso.py'),'R15:'+parts[0]['id']],capture_output=True);assert t.returncode!=0
for p in O.rglob('*.dxf'):assert not ezdxf.readfile(p).audit().has_errors,str(p)
w=openpyxl.load_workbook(O/'TV_R16_muhelyi_regiszter.xlsx',data_only=True);errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e'];assert not errors,errors
assert w['R16_Attekintes']['B11'].value==956 and w['R16_Koltseg']['B23'].value==26212190
old=openpyxl.load_workbook(O/'sources/TV_R15_eredeti_munkafuzet.xlsx',data_only=False);new=openpyxl.load_workbook(O/'TV_R16_muhelyi_regiszter.xlsx',data_only=False)
# All old manual nonformula values outside the intentional summary replacement are retained.
for s in old:
 if s.title=='Attekintes':continue
 for row in s:
  for c in row:
   if c.data_type!='f' and c.value is not None:assert new[s.title][c.coordinate].value==c.value,(s.title,c.coordinate)
master=P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx';a=openpyxl.load_workbook(O/'sources/A902_elo_regiszter_R16_elott.xlsx',data_only=False);b=openpyxl.load_workbook(master,data_only=False);v=openpyxl.load_workbook(master,data_only=True)
allowedrows={31,32,33,35,115};allowedcols={4,8,10,11,12,14,17,23,24,32,33,35,41,42,44}
for s in a:
 for row in s:
  for c in row:
   if s.title=='Elemregiszter' and c.row in allowedrows and c.column in allowedcols:continue
   if c.value is not None and c.data_type!='f':assert b[s.title][c.coordinate].value==c.value,(s.title,c.coordinate,c.value,b[s.title][c.coordinate].value)
assert v['Összesítő']['B7'].value==26212190
assert not [(s.title,c.coordinate) for s in v for row in s for c in row if c.data_type=='e']
book=PdfReader(O/'rajzok/TV_R16_asztalos_tervcsomag.pdf');assert len(book.pages)==121
text='\n'.join(p.extract_text() for p in book.pages);assert all(p['id'] in text for p in parts)
assert len(PdfReader(O/'artwork/INK_W01_R16_1az1_probanyomat.pdf').pages)==7
assert len(PdfReader(O/'cimkek/TV_R16_cimkek_A4.pdf').pages)==11
bl=r('audit/blender_provenance.json');assert len(bl['parts'])==106 and max(p['bbox_error_mm'] for p in bl['parts'])<.001
# Cross-check arithmetic in capacity, lighting and dimensional chains.
assert sum(d['width'] for d in r('data/drawers.json'))==3600
assert abs(sum(c['watts'] for c in r('data/lighting.json')['circuits'])-102.8)<1e-8
shutil.copy2('/tmp/ki16_blender_final.log',O/'audit/blender_run.log')
shutil.copy2(master,O/'A902_Kyoto_Ink_elo_regiszter_R16_atadasi_masolat.xlsx')
report={'revision':'R16','wood_parts':106,'bores':956,'native_volume_mm3':vol,'AutoCAD_saved_solids':215,'hardware_envelopes':109,'art_hatches':145,'AutoCAD_database_errors':0,'blender_bbox_max_mm':max(p['bbox_error_mm'] for p in bl['parts']),'MPR_drafts':len(programs),'MPR_macros':956,'own_MPR_readback':True,'woodWOP_or_machine_test':False,'label_lookups':106,'old_revision_rejected':True,'PDF_pages':[121,7,11],'excel_formula_errors':0,'old_manual_values_preserved':True,'live_project_total':26212190,'live_main_changed_items':['KI-030','KI-031','KI-032','KI-034','KI-114'],'source_image_SHA256':r('data/artwork.json')['sha256'],'not_approved':['survey','MOVENTO template operations','materials and print proof','workshop machine trial','wall and electrical design','coffee table move']}
(O/'audit/release_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
(O/'audit/master_preserved.json').write_text(json.dumps({'before':{'path':'sources/A902_elo_regiszter_R16_elott.xlsx','sha256':hashlib.sha256((O/'sources/A902_elo_regiszter_R16_elott.xlsx').read_bytes()).hexdigest()},'after':{'path':str(master),'sha256':hashlib.sha256(master.read_bytes()).hexdigest()},'all_unrelated_manual_values_preserved':True},indent=2))
manifest=[]
for f in sorted(O.rglob('*')):
 if not f.is_file() or f.name=='file_manifest.json' or f.name.endswith(('.inspect.ndjson','.blend1','.bak','.dwl','.dwl2')):continue
 manifest.append({'path':str(f.relative_to(O)),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(O/'data/file_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
zpath=P/'A902_TV_R16_bronz_tus_fiokok.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
 for x in manifest:z.write(O/x['path'],'R16_feluletek_fiokok/'+x['path'])
 z.write(O/'data/file_manifest.json','R16_feluletek_fiokok/data/file_manifest.json')
 for f in list((P/'scripts').glob('r16_*.py'))+[P/'r16_workbook.mjs']:z.write(f,'rebuild/'+f.name)
with zipfile.ZipFile(zpath) as z:assert z.testzip() is None
print(json.dumps({'files':len(manifest),'zip_MB':zpath.stat().st_size/1e6,'checks':report}))
