from pathlib import Path
import json,re,hashlib,zipfile,shutil,xml.etree.ElementTree as ET,sqlite3,subprocess,sys,tempfile,csv
import openpyxl,ezdxf
from pypdf import PdfReader
P=Path(__file__).resolve().parents[1];O=P/'R15_muhely'
def read(n):return json.loads((O/n).read_text())
parts=read('data/parts.json');holes=read('data/holes.json');programs=read('data/programs.json')
shutil.copy2('/tmp/ki15coordverify.log',O/'audit/autocad_coordination_readback.log')
log=(O/'audit/ki15verify01.log').read_text();vol=float(re.search(r'Volume:\s+([0-9.]+)',log).group(1));expected=read('audit/design_checks.json')['expected_net_volume_mm3']
assert '(KI_R15_SOLIDS 75)' in log and 'Total errors found 0 fixed 0' in log and abs(vol-expected)<.001
assert all(p['id'] in log for p in parts)
c=(O/'audit/autocad_coordination_readback.log').read_text();assert '(R15_ALL_SOLIDS 188)' in c and '(R15_HW_SOLIDS 113)' in c and 'Total errors found 0 fixed 0' in c
# Print-layout metadata: cell values/formulas remain untouched.
f=O/'TV_R15_muhelyi_regiszter.xlsx';ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda s:'{'+ns+'}'+s;ET.register_namespace('',ns)
with zipfile.ZipFile(f) as z:items={n:z.read(n) for n in z.namelist()}
s=ET.fromstring(items['xl/worksheets/sheet12.xml'])
for n in ['pageMargins','pageSetup','rowBreaks']:
 e=s.find(q(n))
 if e is not None:s.remove(e)
ET.SubElement(s,q('pageMargins'),{'left':str(5/25.4),'right':str(5/25.4),'top':str(23.5/25.4),'bottom':str(23.5/25.4),'header':'0','footer':'0'})
ET.SubElement(s,q('pageSetup'),{'paperSize':'9','orientation':'portrait','scale':'100'})
b=ET.SubElement(s,q('rowBreaks'),{'count':'7','manualBreakCount':'7'})
for r in range(50,351,50):ET.SubElement(b,q('brk'),{'id':str(r),'min':'0','max':'16383','man':'1'})
items['xl/worksheets/sheet12.xml']=ET.tostring(s,encoding='utf-8',xml_declaration=True)
w=ET.fromstring(items['xl/workbook.xml']);dn=w.find(q('definedNames'))
if dn is None:dn=ET.Element(q('definedNames'));w.insert(list(w).index(w.find(q('sheets')))+1,dn)
for el in list(dn):
 if el.get('name')=='_xlnm.Print_Area' and el.get('localSheetId')=='11':dn.remove(el)
ET.SubElement(dn,q('definedName'),{'name':'_xlnm.Print_Area','localSheetId':'11'}).text="'Cimkek'!$A$1:$I$380"
items['xl/workbook.xml']=ET.tostring(w,encoding='utf-8',xml_declaration=True)
with zipfile.ZipFile(f,'w',zipfile.ZIP_DEFLATED) as z:
 for n,v in items.items():z.writestr(n,v)
a=openpyxl.load_workbook(f,data_only=False);v=openpyxl.load_workbook(f,data_only=True)
errs=[[sh.title,x.coordinate,x.value] for sh in v for row in sh for x in row if x.data_type=='e'];assert not errs,errs
assert [v['Attekintes'][f'B{i}'].value for i in [2,3,4,5]]==[75,826,73,826]
assert v['Koltseg']['I23'].value==802800 and v['Koltseg']['B31'].value==25161190
old=read('audit/R14_excel_snapshot.json');manual={r[0]:r for r in old['sheets']['Darabok'][5:] if r[0]}
for i,p in enumerate(parts,6):
 for j in list(range(7,13))+[16,17]:assert a['Darabok'].cell(i,j).value==manual[p['id']][j-1],(p['id'],j)
master=P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx';assert hashlib.sha256(master.read_bytes()).hexdigest()==read('audit/project_budget_readback.json')['sha256']
books=[('rajzok/TV_R15_szerelesi_es_alkatreszrajzok.pdf',103),('pilot/M5_elso_modul_probacsomag.pdf',15),('cimkek/TV_R15_cimkek_A4.pdf',8)]
for name,n in books:assert len(PdfReader(O/name).pages)==n
text='\n'.join(p.extract_text() for p in PdfReader(O/books[0][0]).pages);assert all(p['id'] in text for p in parts)
for root in [O,O/'pilot']:
 ps=json.loads((root/'data/parts.json').read_text());pr=json.loads((root/'data/programs.json').read_text())
 for q0 in pr:
  raw=(root/q0['path']).read_bytes();assert hashlib.sha256(raw).hexdigest()==q0['sha256'];assert b'EN="0"' not in raw and b'??="0"' not in raw
 for p in ps:
  t=subprocess.run([sys.executable,str(root/'integration/cimke_kereso.py'),'R15:'+p['id']],capture_output=True,text=True);assert t.returncode==0,(p['id'],t.stderr)
 t=subprocess.run([sys.executable,str(root/'integration/cimke_kereso.py'),'R14:'+ps[0]['id']],capture_output=True);assert t.returncode!=0
 with (root/'integration/alkatresz_import_UTF8.csv').open(encoding='utf-8-sig') as f:
  rows=list(csv.DictReader(f,delimiter=';'));assert len(rows)==len(ps)
  for r in rows:assert (root/r['drawing']).is_file()
with tempfile.TemporaryDirectory(prefix='ki15-integrity-') as td:
 t=Path(td);(t/'integration').mkdir();shutil.copy2(O/'integration/cimke_kereso.py',t/'integration');shutil.copy2(O/'integration/alkatreszek_R15.sqlite',t/'integration')
 db=sqlite3.connect(t/'integration/alkatreszek_R15.sqlite');row=db.execute('SELECT dxf,programs FROM parts WHERE barcode=?',('R15:'+parts[0]['id'],)).fetchone();db.close()
 files=[row[0]]+json.loads(row[1])
 for name in files:(t/name).parent.mkdir(exist_ok=True,parents=True);shutil.copy2(O/name,t/name)
 target=t/files[-1];target.write_bytes(target.read_bytes()+b'\nCORRUPT_TEST\n');r=subprocess.run([sys.executable,str(t/'integration/cimke_kereso.py'),'R15:'+parts[0]['id']],capture_output=True,text=True);assert r.returncode!=0 and 'Hash eltérés' in r.stderr
for f in O.rglob('*.dxf'):assert not ezdxf.readfile(f).audit().has_errors,str(f)
results=dict(wood_solids=75,hardware_envelopes=113,bores=826,notches=12,net_volume_mm3=vol,analytic_volume_mm3=expected,AutoCAD_database_errors=0,main_MPR=73,pilot_MPR=11,pilot_parts=10,own_macro_readback=True,woodWOP_test=False,machine_test=False,main_label_lookups=75,pilot_label_lookups=10,obsolete_revision_rejected=True,corrupt_file_rejected=True,XLSX_errors=0,R14_manual_fields_preserved=True,master_workbook_unchanged=True,PDF_pages=[103,15,8],door_variant_selected=None,physical_survey=False,physical_hardware_trial=False)
(O/'audit/release_checks.json').write_text(json.dumps(results,ensure_ascii=False,indent=2))
manifest=[]
for f in sorted(O.rglob('*')):
 if not f.is_file() or f.name=='file_manifest.json' or f.name.endswith('.inspect.ndjson') or f.suffix in ['.dwl','.dwl2','.blend1','.bak']:continue
 manifest.append({'path':str(f.relative_to(O)),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(O/'data/file_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
zpath=P/'A902_TV_R15_hat_fejlesztes.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
 for x in manifest:z.write(O/x['path'],'R15_muhely/'+x['path'])
 z.write(O/'data/file_manifest.json','R15_muhely/data/file_manifest.json')
 for f in list((P/'scripts').glob('r15_*.py'))+[P/'r15_workbook.mjs']:z.write(f,'rebuild/'+f.name)
 z.writestr('rebuild/OLVASD_EL.txt','A kész fájlok újragenerálás nélkül használhatók. A generátorok a meglévő execution munkamappát és R12–R14 forrásait, a helyi AutoCAD/Blender és Python/Node környezetet használják. Ez nem önálló telepítő. Bizonytalan kimenetelű CAD-módosítást ne ismételj automatikusan.')
with zipfile.ZipFile(zpath) as z:assert z.testzip() is None
print(json.dumps({'files':len(manifest),'zip_bytes':zpath.stat().st_size,'checks':results}))
