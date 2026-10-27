from pathlib import Path
import json,hashlib,zipfile,re,csv,subprocess
import openpyxl,pdfplumber
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
P=Path(__file__).resolve().parents[1];O=P/'R05';root=P.parents[1]
def read(f):return json.loads((P/f).read_text())
def write(f,s):(O/f).write_text(s)
C=read('R05/data/cost_R05.json');base=read('R05/sources/workbook_before_R05.json');touch={(s,a) for s,a,v in read('R05/qa/touched_cells.json')};w=openpyxl.load_workbook(P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx');v=openpyxl.load_workbook(P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx',data_only=True)
norm=lambda x:None if x=='' else x;diff=[]
for s,data in base.items():
 for ri,row in enumerate(data,1):
  for ci,old in enumerate(row,1):
   addr=get_column_letter(ci)+str(ri);new=w[s].cell(ri,ci).value
   if norm(old)!=norm(new) and (s,addr) not in touch:diff.append([s,addr,old,new])
assert not diff,diff[:8]
assert v['Összesítő']['B7'].value==C['main'] and v['Összesítő']['B8'].value==C['economy']
assert sum(v['CP1a eltérések'].cell(r,3).value or 0 for r in range(2,161))==20083350
assert sum(v['CP3 eltérések'].cell(r,3).value or 0 for r in range(2,174))==23253840
assert len(PdfReader(O/'owner/A902_R05_tulajdonosi_utmutato_5oldal.pdf').pages)==5
write('qa/workbook_preservation.json',json.dumps({'untargeted_cell_changes':diff,'preserved_existing_sheets':20,'final_sheets':len(w.sheetnames),'rows':174,'calculated_main_readback':C['main'],'calculated_economy_readback':C['economy'],'semantic_history_repairs':read('R05/qa/history_reconciliation.json')},ensure_ascii=False,indent=2))
ref=Path('CP3/reviews/20261027_csomag_es_mediafal/A902_KyotoInk_V01_KI-CP3-R03_VISUAL_RESCUE_v3_CONDITIONAL.png')
changes=[
 ['R05-01','Médiafal tájolás','Régi torony Y7,30–8,00m','ToronyY4,70–5,40m;panelY5,40–8,00m;TV+700mm','Elfogadott kép szerinti oldal; kiállások és rögzítés felmérendő'],
 ['R05-02','W02 nyomat','150×220cm keretezett kép árazása','2100×2880mm keret nélküli mező;8m² rendelési keret','+121000Ft a korábbi R04 összeghez;2nyomtatott mező'],
 ['R05-03','Költség és előzmények','R04 külön forgatókönyv, CP3 élő összeg','174sorosR05 élő regiszter; CP1a hibás hivatkozása javítva;CP3előzmény rögzítve','18M cél;2M tartalék0felhasználás; nincs megrendelés'],
 ['R05-04','Gazdaságos változat','Fiók-, fejvégszín- és kültéri anyagelhagyások is szerepeltek','253000Ft három becsült csere; tárolókapacitás változatlan','SíkabbW01,vliesW02,fixdugalj; egyik sem jóváhagyott'],
 ['R05-05','Anyagcsalád','Általános kréta/grafit megjelölés','10anyagkód33egyedi szerkezethez rendelve','Több konkrét gyártói dekor/hordozó; bevonat/kárpitSKU és fizikai minták nyitottak'],
 ['R05-06','W01fény','Általános LEDcsomag','Loox5 LED2074 / 833.74.308; 3 ág; 2 táp; 79,008 W','A költség150000Ft becslés; Hue,profil,tápSKU és hőpróba nyitott'],
 ['R05-07','Grafikai állomány','Szemléltető fekete jel a rajzon','Elfogadott eredeti kép+3szerkeszthető nyomdai kiosztási sablon','A sablon NEM kész tusgrafika; nincs hamis nyomdai jóváhagyás'],
 ['R05-08','Renderátadás','CP3régi modellek/rajzok/promptok','28egyedi nézet+21funkcionális prompt,R05modell- és regiszterhash','V01új natív kontroll;27régi kamerakontroll csak archív perspektívareferencia']]
write('data/change_log.json',json.dumps(changes,ensure_ascii=False,indent=2))
write('OLVASD_EL_R05.txt',f'''KYOTO INK / KI-R05 / 2026-10-27

LEGÚJABB MUNKAVÁLTOZAT, NEM LEZÁRT MEGVALÓSÍTÁS
Az R04 mentésből folytatva. A korábbi rajzokat és kreatív forrásokat megőriztük.
18 000 000 Ft célkeret; 2 000 000 Ft külön tartalék, érintetlen.
A 23,79M Ft elfogadható tervezési jelzés nem rendelési vagy keretemelési engedély.

NYISD EZEKET
1. owner/A902_R05_tulajdonosi_utmutato_5oldal.pdf — öt A4 oldal, aktuális költség, tárolás, berendezési vázlat és gyorsítási lehetőség.
2. drawings/A902_R05_aktualis_butorrajzok.pdf — 71 A3 lap,33egyedi szerkezet, médiafal/W02/villamos részlet, mindkét tároló.13érdemben frissített lap, a többi érvényes R04lap változatlanul beemelve. Az egyedi PDF/SVG eredetét data/drawing_index.json mutatja; nincs két, egyformán aktuális kiosztás.
3. materials/A902_R05_anyag_es_ajanlatkeresi_csomag.pdf — 14A4oldal:10anyagcsalád, költségösszehasonlítás és3szakági ajánlatkérési tartalom.
4. ../outputs/A902_Kyoto_Ink_elo_regiszter.xlsx — ugyanaz az élő Excel;174sor,23munkalap. Képletek újraszámolva, mentésből visszaolvasva. Nem keletkezett második élő költségvetés.
5. blender/A902_Kyoto_Ink_R05.blend — külön mentett és újranyitott műszaki munkamodell. Két natív kontrollkép. Egyszerűsített késztermékek, nem termékazonos végleges látvány.
6. render_handoff/START_RENDER_OPERATOR_EN.txt — egy másolható indítóprompt;28külön nézet és21funkcionális állapot teljes angol promptja. Tényleges fájlokat is csatolni kell, nem elég a fájlnevet bemásolni.

KÖLTSÉG
Fő:{C['main']:,}Ft. Takarékos összehasonlítás:{C['economy']:,}Ft.
Tételes becsült megtakarítás:{C['saving']:,}Ft. A takarékos változat is{C['economy_gap']:,}Ft-tal az18M felett van.
Becsült sáv:{C['low']:,}–{C['high']:,}Ft; nem valószínűségi garancia.
Ellenőrzött webshopárú sorok:{C['verified']:,}Ft. Írásos beszállítói ajánlat:0Ft. Becsült rész:{C['estimated']:,}Ft.
Nem hagytuk el a jakuzzit, a második fagyasztót vagy a kért fiókokat. Fix kivitelezői ár nincs.

MI JAVULT KONKRÉTAN?
A célképpel összehasonlított Blender-próbán a médiafal tornya a rossz végén állt. Átkerült a kiválasztott kép szerinti oldalra, a panel és TV együtt eltolva. Az alsó700mm-es front most a torony alatt van.
A W02teljes falmező anyaga, rendelési mennyisége és hordozója most ugyanahhoz a költséghez kötődik. A régi kisebb keretezett kép már nem számít bele.
A CP1a összehasonlító oszlop rossz helyre hivatkozó képleteit az eredeti mentett mennyiség és egységár alapján javítottuk:20 083 350Ft. A CP3történeti érték23 253 840Ft; nem változik többé az aktuális árakkal.

MI ALKALMAS AJÁNLATKÉRÉSRE?
A méretezett tervezési kiosztás, anyagjelöltek, mennyiségek, alternatív technológiák és a kért ajánlati bontás. Felméréshez és gyártói ellenőrzéshez világos alap, nem kész gyártási méretkiadás. A mintázási és nyitott műszaki feladatokat az ajánlatban külön kell szerepeltetni.

MI MARADT NYITVA?
Eredeti nagyfelbontású/vektoros tusmaster; fizikai szín-, kárpit-, felület- és fénypróba; pontos bevonat/pánt/táp/profilSKU; beszállítói ár és magyar szállítás; gyártói mozgás és helyszíni méretek.
A teljes lakás használati ütközései nincsenek lezárva: konyhai840mmfolyosó és nyitott mosogatógép;05.3szék/ágy; master ágyfiók/éjjeli; liftkinematika; tárolóajtók és bicikli kivétele; küszöbszinteltérés.
Az SVGgrafikai sablon üres nyomat-rétege szándékos: a kiosztás kész, a tusgrafikát nem helyettesítettük kitalált új motívummal. A késztermékek SKUja önmagában nem bizonyítja, hogy a3Dmodell alakja pontos.

KÖVETKEZŐ ELLENŐRZÉSI PONT
Az első V01natív próba akkor léphet termékazonos látványfázisba, ha a Lennon/Eamy/Toni/Somar valós alakja, a médiafal eredeti tusgrafikája és anyagmintája a modellbe került. Ezzel párhuzamosan felmérés és egységes ajánlati kör zárja a méretet és árat. Új kép nem minősül véglegesnek ezek helyett.
Az ajánlatkérések előkészítve, nem elküldve. Rendelés és kivitelezői kötelezettségvállalás nem történt.
''')
write('evidence/forrasok_20261027.txt','''R05 műszaki forrásellenőrzés /2026-10-27
Ár és műszaki adat nem ugyanaz. A felsorolt anyagok többségéhez NINCS hazai árajánlat.

EGGER U702PM: meleg szürke, matt/ujjlenyomatcsökkentett felület; a szín csak eredeti mintán dönthető el.
https://www.egger.com/en/furniture-interior-design/decors/U702_PM?country=US
U702/U961PM MDF termékcsalád:
https://downloads.egger.com/static/uk/FY_EGGER_PerfectSense_Portfolio_A4_Screen.pdf
F274ST9 LightConcrete dekor szerepel a2025–26gyártói választékban; ez nem hazai munkalap-SKU vagy raktárigazolás.
https://www.egger.com/get_download/1ed4150d-4c94-4717-a4d9-81d28cacddf0/Product_Availability_Industrial_Decors_English.pdf?country=NL

Vescom Bely+Print1040.01:134cm teljes,133cm nyomtatható szélesség;350g/m²;vinil/viszkóz-poliészter. Egyedi nyomat jelölt.
https://vescom.com/en/products/belyprint
https://vescom.com/customise/wallcovering-print

Häfele LED2074/833.74.308:12V,2700K,9,6W/m,CRI>90;4m ajánlott maximális ághossz. Ezért PS-A két külön ágat kap.
https://files.hafele.co.uk/catalogfiles/www/25-1010.pdf
Gledopto GL-C-009P:12–54V,5A/csatorna,10Aösszesen. Hue-integrációt nem tekintjük helyben teszteltnek.
https://gledopto.com/h-pd-134.html

Steadyrack ClassicRoad:89,99EUR megjelenített ár;20–29col kerék,max2,1col gumi,35kg,sárvédő nélkül. Magyar szállítás és a valós bicikli nincs igazolva.45000Ft beszerzési keret+25000Ft rögzítés maradt a főösszegben.
https://eu.steadyrack.com/products/classic-road-bike-rack

IKEA MALM180×200/004.048.12:239900Ft,210×195cm külső méret. Megvizsgált, NEM kiválasztott alternatíva; NORDLIhoz képest nagyobb külső méret és fejvég, ezért nem oldottuk meg vele igazolatlanul az ágyfiók-ütközést.
https://www.ikea.com/hu/hu/p/malm-agynemutartos-agy-feher-00404812/

Blum TANDEM560H: gyártói30kg teljes kihúzás,11–16mm fadoboz,500mm változat560H5000B. A meglévő terv általános dobozmérete nem automatikusan gyártói szabásméret; illesztési rajz szükséges.
https://publications.blum.com/2024/catalogue/th/434/

KaveHomeEamy/Somar friss közvetlen megnyitása429hibát adott. A korábbi dátumú SKU/ár maradt; nem jelöltük ma újra ellenőrzött árnak. Gyártói3Dmodellt nem sikerült igazoltan beszerezni. A korábbiCP3IKEAtermékfotók megtartva; a PAXfotó továbbra sem a kiválasztott egyedi gardrób.
''')
readme=P/'README.md';s=readme.read_text();header='## R05 — aktuális ár-, anyag- és render-egyeztetés (2026-10-27)\n\nBelépő: [R05/OLVASD_EL_R05.txt](R05/OLVASD_EL_R05.txt). Élő Excel: **23 912 890 Ft** fő / **23 659 890 Ft** takarékos összehasonlítás.18M cél,2M érintetlen tartalék.71A3 aktuális rajzlap összeállítva,14oldalas anyag/RFQcsomag,5oldalas tulajdonosi PDF, külön újranyitott Blender-modell és49nézeti/funkcionális prompt. **Feltételes árazási és egyeztetési kiadás, nem kész megvalósítás vagy gyártási jóváhagyás.**\n\n'
if not s.startswith('## R05'):readme.write_text(header+s)
qa={'owner_pages':5,'current_drawing_pages':len(PdfReader(O/'drawings/A902_R05_aktualis_butorrajzok.pdf').pages),'material_Rfq_pages':14,'formula_errors':read('R05/qa/excel_checks.json'),'preservation':read('R05/qa/workbook_preservation.json'),'pdf_bounds':read('R05/qa/pdf_text_bounds.json'),'blender_reopen':read('R05/qa/blender_reopen.json'),'handoff':read('R05/qa/render_handoff_checks.json'),'visual_review':['owner five-page layout','W01 target beside technical elevation','W02 print and mounting detail','native V01 before/after tower correction','updated media wall orthographic control','Excel main summary and material/change/capacity ranges'],'final_product_render_approved':False}
write('qa/ELLENORZESI_NYUGTA.json',json.dumps(qa,ensure_ascii=False,indent=2))
# Upload only the current files and actual binaries; temporary autosaves/previews excluded.
files=[f for f in O.rglob('*') if f.is_file() and 'previews' not in f.parts and f.name not in ['DELIVERY_MANIFEST.json','GITHUB_VISSZAOLVASAS.json','blender_source_inventory.json'] and f.suffix not in ['.blend1','.pyc']]
files+=list((P/'scripts').glob('r05_*.py'))+[P/'r05_register.mjs',readme,P/'outputs/A902_Kyoto_Ink_elo_regiszter.xlsx']
rows=[]
for f in sorted(set(files)):
 b=f.read_bytes();rows.append({'path':str(f.relative_to(root)),'local':str(f),'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'git_sha':hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()})
manifest=O/'DELIVERY_MANIFEST.json';manifest.write_text(json.dumps([{k:v for k,v in r.items() if k!='local'} for r in rows],ensure_ascii=False,indent=2));b=manifest.read_bytes();rows.append({'path':str(manifest.relative_to(root)),'local':str(manifest),'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest(),'git_sha':hashlib.sha1(b'blob '+str(len(b)).encode()+b'\0'+b).hexdigest()});Path('/tmp/r05_uploadfiles.json').write_text(json.dumps(rows,ensure_ascii=False))
# ZIP is self-contained for every render attachment and retained drawing source.
zfiles={Path(r['local']) for r in rows};m=read('R05/render_handoff/manifest_R05.json')
for vw in m['views']:
 for a in vw['attachments']:zfiles.add(P/a['path'])
for a in read('R05/data/drawing_index.json'):
 f=P/a['path'];zfiles.add(f)
 if f.with_suffix('.svg').exists():zfiles.add(f.with_suffix('.svg'))
for f in [P/'data/render_manifest.json',P/'CP3/data/tulajdonosi_dontes.json',P/'sources/TELJES_MESTERTERV_MEGRENDELOI_BRIEF.txt']:
 if f.exists():zfiles.add(f)
zpath=P/'outputs/A902_R05_egyeztetesi_atadas.zip'
with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for f in sorted(zfiles):z.write(f,str(f.relative_to(P)))
with zipfile.ZipFile(zpath) as z:assert z.testzip() is None
print(json.dumps({'upload_files':len(rows),'upload_MB':round(sum(r['bytes'] for r in rows)/1e6,2),'zip_files':len(zfiles),'zip_MB':round(zpath.stat().st_size/1e6,2),'unrelated_changes':len(diff),'owner_pages':5,'drawing_pages':qa['current_drawing_pages']}))
